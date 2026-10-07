"""FastAPI app for A13 Career Planning Agent.

Run:
    uvicorn app.backend.main:app --reload --port 8000
"""
from __future__ import annotations
import hashlib
import io
import json
import random
import re
import secrets
import time
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException, Request, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from starlette.responses import JSONResponse as _JSONResponse
from starlette.types import ASGIApp, Receive, Scope, Send
from pydantic import BaseModel

from . import llm, matching
from .store import store

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "app" / "data"

app = FastAPI(title="A13 Career Planning Agent")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


import traceback as _tb

@app.exception_handler(Exception)
async def _global_exc(request: Request, exc: Exception):
    _tb.print_exc()          # 完整堆栈打到终端
    return _JSONResponse({"detail": str(exc)}, status_code=500)

# ---------- session / auth ----------
_sessions: dict[str, str] = {"demo": "demo"}  # token -> username; "demo" always valid
_sms_codes: dict[str, tuple[str, float]] = {}  # phone -> (code, expire_ts)

_PUBLIC_PREFIXES = ("/api/login", "/api/register", "/api/sms/", "/docs", "/openapi.json")


def _hash_pw(pw: str) -> str:
    return hashlib.sha256(pw.encode()).hexdigest()


class AuthMiddleware:
    """Pure ASGI middleware — does NOT read request body (avoids BaseHTTPMiddleware bug)."""

    def __init__(self, app: ASGIApp):
        self.app = app

    async def __call__(self, scope: Scope, receive: Receive, send: Send):
        if scope["type"] != "http":
            await self.app(scope, receive, send)
            return

        path: str = scope["path"]
        method: str = scope.get("method", "GET")

        # Only API routes require bearer authentication.  The production build
        # serves the React application from this same FastAPI process, so page
        # routes such as /login and /report must remain publicly reachable.
        if (not path.startswith("/api")) or method == "OPTIONS" or any(path.startswith(p) for p in _PUBLIC_PREFIXES):
            await self.app(scope, receive, send)
            return

        # Extract token from headers (ASGI headers are list of [name, value] byte-pairs)
        token = ""
        for hname, hval in scope.get("headers", []):
            if hname == b"authorization":
                token = hval.decode().removeprefix("Bearer ").strip()
                break

        if not token or token not in _sessions:
            resp = _JSONResponse({"detail": "未登录或 token 已失效"}, status_code=401)
            await resp(scope, receive, send)
            return

        # Inject username into scope.state so Request.state.username works
        scope.setdefault("state", {})["username"] = _sessions[token]
        await self.app(scope, receive, send)


app.add_middleware(AuthMiddleware)


def _current_user(request: Request) -> str:
    return getattr(request.state, "username", "demo")


def _load(name: str) -> Any:
    return json.loads((DATA / name).read_text(encoding="utf-8"))


POSITIONS = _load("positions.json")
GRAPH = _load("graph.json")
POS_BY_ID = {p["id"]: p for p in POSITIONS}

# 知乎工作体验数据（丰富"典型的一天"）
_ZHIHU_PATH = DATA / "zhihu_experience.json"
ZHIHU_EXP: dict = json.loads(_ZHIHU_PATH.read_text(encoding="utf-8")) if _ZHIHU_PATH.exists() else {}

# 职业规划智慧库（来自真实咨询案例）
_WISDOM_PATH = DATA / "career_wisdom.json"
CAREER_WISDOM: dict = json.loads(_WISDOM_PATH.read_text(encoding="utf-8")) if _WISDOM_PATH.exists() else {}
_CONFUSION_PATTERNS: list[dict] = CAREER_WISDOM.get("confusion_patterns", [])
_UNIVERSAL_PRINCIPLES: list[dict] = CAREER_WISDOM.get("universal_principles", [])
_QUOTES_BANK: list[str] = CAREER_WISDOM.get("quotes_bank", [])

# 张雪峰方法论语录库
_ZXF_PATH = DATA / "zxf_wisdom.json"
ZXF_WISDOM: list[dict] = json.loads(_ZXF_PATH.read_text(encoding="utf-8")) if _ZXF_PATH.exists() else []


def _zxf_by_tags(tags: list[str], limit: int = 2) -> list[dict]:
    """根据标签从张雪峰语录库中匹配最相关的条目"""
    if not ZXF_WISDOM or not tags:
        return []
    tag_set = {t.lower() for t in tags}
    scored: list[tuple[int, dict]] = []
    for q in ZXF_WISDOM:
        qtags = {t.lower() for t in q.get("tags", [])}
        overlap = len(tag_set & qtags)
        if overlap > 0:
            scored.append((overlap, q))
    scored.sort(key=lambda x: -x[0])
    return [q for _, q in scored[:limit]]


def _zxf_by_major(major: str, limit: int = 2) -> list[dict]:
    """根据专业从张雪峰语录库中匹配"""
    if not ZXF_WISDOM or not major:
        return []
    results: list[dict] = []
    for q in ZXF_WISDOM:
        for m in q.get("related_majors", []):
            if m in major or major in m:
                results.append(q)
                break
        if len(results) >= limit:
            break
    return results


def _zxf_by_job(job_name: str, limit: int = 2) -> list[dict]:
    """根据岗位名从张雪峰语录库中匹配"""
    if not ZXF_WISDOM or not job_name:
        return []
    # 岗位名 → 可能的标签
    job_tags: list[str] = []
    tag_map = {
        "金融": ["金融", "银行", "经济"],
        "计算机": ["计算机", "互联网", "AI冲击"],
        "医": ["医学", "三甲医院"],
        "法": ["法学"],
        "教": ["教育", "师范", "教培"],
        "设计": ["设计", "艺术"],
        "销售": ["销售", "起薪"],
        "公务": ["考公", "编制", "体制内"],
        "运营": ["互联网", "自媒体"],
        "数据": ["数据科学", "统计学"],
        "产品": ["互联网"],
        "会计": ["会计"],
        "翻译": ["外语", "小语种"],
        "建筑": ["建筑学", "土木工程"],
        "化工": ["生化环材", "天坑专业"],
        "机械": ["工科"],
        "电气": ["电气工程", "国家电网"],
        "新闻": ["新闻学", "自媒体"],
    }
    name_lower = job_name.lower()
    for kw, tags in tag_map.items():
        if kw in name_lower:
            job_tags.extend(tags)
    if not job_tags:
        job_tags = ["就业", "职业发展"]
    return _zxf_by_tags(job_tags, limit)


def _detect_confusion(prof: dict, assess: dict | None = None) -> list[dict]:
    """根据用户档案特征，匹配可能的困惑模式。"""
    hits: list[dict] = []
    assess = assess or {}
    intent = prof.get("intent") or {}
    internships = prof.get("internships") or []
    work = prof.get("work_experiences") or []
    skills = prof.get("skills") or []
    all_jobs = list(internships) + list(work)

    # 频繁跳槽
    if len(all_jobs) >= 3:
        industries = set()
        for e in all_jobs:
            if isinstance(e, dict) and e.get("position"):
                industries.add(e["position"])
        if len(industries) >= 3:
            hits.append(_get_pattern("job_hopping"))

    # 躺平停滞（无经历 + 有学历）
    if not all_jobs and not skills:
        hits.append(_get_pattern("stagnation"))

    # 求职受挫（有经历但技能和岗位不匹配多）
    # 毕业前混沌（在校 + 无实习）
    edu = prof.get("education") or []
    if edu and not internships and not work:
        hits.append(_get_pattern("pre_graduation_chaos"))

    # 横向比较（通用，几乎每个人都有）
    hits.append(_get_pattern("comparison_anxiety"))

    return [h for h in hits if h]


def _get_pattern(pid: str) -> dict | None:
    for p in _CONFUSION_PATTERNS:
        if p.get("id") == pid:
            return p
    return None


def _pick_quotes(n: int = 2) -> list[str]:
    if not _QUOTES_BANK:
        return []
    return random.sample(_QUOTES_BANK, min(n, len(_QUOTES_BANK)))


# ---------- helpers ----------
def progress(user_state: dict) -> dict:
    profile = user_state.get("profile") or {}
    completeness = _completeness(profile)
    profile_started = bool(profile.get("basic", {}).get("name"))
    assess = user_state.get("assessments") or {}
    assess_done = all(assess.get(k) for k in ("holland", "multi_intel", "values", "mbti"))
    report_done = bool(user_state.get("report"))
    return {
        "resume": "已完善" if completeness >= 100 else "待完善",
        "self": "已完善" if assess_done else "待完善",
        "plan": "已完成" if report_done else ("进行中" if profile_started else "未开始"),
    }


# ---------- auth / user ----------
class LoginIn(BaseModel):
    username: str
    password: str


class RegisterIn(BaseModel):
    username: str
    password: str


class SmsSendIn(BaseModel):
    phone: str


class SmsLoginIn(BaseModel):
    phone: str
    code: str


@app.post("/api/register")
def register(body: RegisterIn):
    if not body.username or not body.password:
        raise HTTPException(400, "用户名和密码不能为空")
    ok = store.add_user(body.username, _hash_pw(body.password))
    if not ok:
        raise HTTPException(400, "用户名已存在")
    token = secrets.token_hex(16)
    _sessions[token] = body.username
    return {"token": token, "username": body.username}


@app.post("/api/login")
def login(body: LoginIn):
    user = store.get_user(body.username)
    if not user or user["password_hash"] != _hash_pw(body.password):
        raise HTTPException(401, "用户名或密码错误")
    token = secrets.token_hex(16)
    _sessions[token] = body.username
    return {"token": token, "username": body.username}


@app.post("/api/sms/send")
def sms_send(body: SmsSendIn):
    if not body.phone or len(body.phone) != 11:
        raise HTTPException(400, "手机号格式不正确")
    code = f"{random.randint(0, 999999):06d}"
    _sms_codes[body.phone] = (code, time.time() + 300)
    print(f"[SMS] 验证码 → {body.phone}: {code}")
    return {"ok": True}


@app.post("/api/sms/login")
def sms_login(body: SmsLoginIn):
    entry = _sms_codes.get(body.phone)
    if not entry:
        raise HTTPException(401, "请先发送验证码")
    saved_code, expire = entry
    if time.time() > expire:
        del _sms_codes[body.phone]
        raise HTTPException(401, "验证码已过期")
    if body.code != saved_code:
        raise HTTPException(401, "验证码错误")
    del _sms_codes[body.phone]
    # find or create user by phone
    username = store.find_user_by_phone(body.phone)
    if not username:
        username = f"phone_{body.phone}"
        store.add_user(username, _hash_pw(secrets.token_hex(8)), phone=body.phone)
    token = secrets.token_hex(16)
    _sessions[token] = username
    return {"token": token, "username": username}


@app.get("/api/me")
def me(request: Request):
    user = _current_user(request)
    s = store.get(user)
    return {"avatar": s.get("avatar"), "progress": progress(s), "username": user}


class AvatarIn(BaseModel):
    avatar: str


@app.post("/api/avatar")
def set_avatar(body: AvatarIn, request: Request):
    store.patch("avatar", body.avatar, user=_current_user(request))
    return {"ok": True}


# ---------- resume / profile ----------
RESUME_SKILL_LEXICON = [
    "Java", "Python", "C++", "C语言", "JavaScript", "TypeScript", "Go", "Rust",
    "PHP", "Kotlin", "Swift", "Scala", "R语言", "MATLAB", "SQL",
    "HTML", "CSS", "Vue", "React", "Angular", "小程序", "Webpack", "Vite",
    "SpringBoot", "Spring", "MyBatis", "Django", "Flask", "FastAPI", "Node.js",
    ".NET", "MySQL", "Redis", "MongoDB", "Oracle", "PostgreSQL", "Kafka",
    "Hadoop", "Spark", "Flink", "TensorFlow", "PyTorch", "OpenCV", "NLP",
    "机器学习", "深度学习", "数据分析", "Excel", "SPSS", "Tableau", "PowerBI",
    "Linux", "Docker", "K8s", "Kubernetes", "Git", "Jenkins", "CI/CD", "AWS",
    "阿里云", "腾讯云", "JUnit", "Selenium", "Jmeter", "Postman",
]

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"1[3-9]\d{9}")


def _extract_text(content: bytes, filename: str) -> str:
    """Extract plain text from PDF / DOCX / TXT. Sniffs magic bytes as well
    as filename extension so uploads without a correct extension still work.
    For PDFs we try PyMuPDF first (usually better for Chinese), then
    pdfplumber as a fallback, and keep whichever produced more text.
    """
    name = (filename or "").lower()
    is_pdf = name.endswith(".pdf") or content[:5] == b"%PDF-"
    is_zip = content[:2] == b"PK"
    is_docx = name.endswith((".doc", ".docx")) or is_zip

    if is_pdf:
        candidates: list[str] = []
        # PyMuPDF (fitz)
        try:
            import fitz  # type: ignore
            doc = fitz.open(stream=content, filetype="pdf")
            txt = "\n".join(page.get_text() for page in doc)
            doc.close()
            if txt.strip():
                candidates.append(txt)
        except Exception as exc:
            print(f"[resume] PyMuPDF failed: {exc}")
        # pdfplumber
        try:
            import pdfplumber  # type: ignore
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                txt = "\n".join((p.extract_text() or "") for p in pdf.pages)
            if txt.strip():
                candidates.append(txt)
        except Exception as exc:
            print(f"[resume] pdfplumber failed: {exc}")
        if candidates:
            # Prefer whichever extracted more characters
            return max(candidates, key=len)
    if is_docx:
        try:
            import docx  # type: ignore
            d = docx.Document(io.BytesIO(content))
            parts = [p.text for p in d.paragraphs]
            for tbl in d.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
            txt = "\n".join(parts)
            if txt.strip():
                return txt
        except Exception as exc:
            print(f"[resume] python-docx failed: {exc}")
    # Plain-text fallback
    try:
        return content.decode("utf-8", errors="ignore")
    except Exception:
        return ""


SECTION_HEADERS = {
    "education": ["教育经历", "教育背景", "学习经历", "Education"],
    "internships": ["实习经历", "实习经验", "Internship"],
    "work_experiences": ["工作经历", "工作经验", "职业经历", "Work Experience"],
    "projects": ["项目经历", "项目经验", "Project"],
    "competitions": ["竞赛", "比赛经历", "获奖"],
    "certificates": ["证书", "技能证书", "资格证书"],
    "self_eval": ["自我评价", "个人评价", "Self-Assessment"],
    "skills_sec": ["专业技能", "技能特长", "技术栈", "Skills"],
}

DATE_RANGE_RE = re.compile(
    r"(\d{4}[./\-年]\s*\d{1,2})\s*[-—~至到]{1,2}\s*"
    r"(\d{4}[./\-年]\s*\d{1,2}|至今|现在|present)",
    re.IGNORECASE,
)
SINGLE_DATE_RE = re.compile(r"(\d{4})[./\-年]\s*(\d{1,2})")


def _split_sections(text: str) -> dict[str, str]:
    """Locate known section headers and return text slices per section."""
    # flatten all header aliases -> canonical key
    header_to_key: list[tuple[str, str]] = []
    for key, aliases in SECTION_HEADERS.items():
        for a in aliases:
            header_to_key.append((a, key))
    # Find positions
    hits: list[tuple[int, str, str]] = []
    for alias, key in header_to_key:
        idx = 0
        while True:
            pos = text.find(alias, idx)
            if pos == -1:
                break
            hits.append((pos, key, alias))
            idx = pos + len(alias)
    hits.sort()
    result: dict[str, str] = {}
    for i, (pos, key, alias) in enumerate(hits):
        start = pos + len(alias)
        end = hits[i + 1][0] if i + 1 < len(hits) else len(text)
        chunk = text[start:end].strip(" :：\n\r\t")
        # keep the longest chunk per key
        if len(chunk) > len(result.get(key, "")):
            result[key] = chunk
    return result


def _norm_month(s: str) -> str:
    s = s.strip()
    m = SINGLE_DATE_RE.search(s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}"
    return s


def _parse_education_block(chunk: str) -> list[dict]:
    """Each paragraph (separated by a date range or blank line) is one entry."""
    items: list[dict] = []
    # split by blank lines first
    blocks = [b.strip() for b in re.split(r"\n\s*\n", chunk) if b.strip()]
    if len(blocks) <= 1:
        # fallback: split by date ranges
        blocks = []
        last = 0
        for m in DATE_RANGE_RE.finditer(chunk):
            if m.start() > last:
                blocks.append(chunk[last:m.start()].strip())
            last = m.start()
        if last < len(chunk):
            blocks.append(chunk[last:].strip())
        blocks = [b for b in blocks if b]
    for block in blocks[:5]:
        item = {
            "period_start": "", "period_end": "", "school": "",
            "degree": "", "college": "", "major": "",
            "direction": "", "advisor": "",
        }
        m = DATE_RANGE_RE.search(block)
        if m:
            item["period_start"] = _norm_month(m.group(1))
            end = m.group(2)
            item["period_end"] = "" if end in ("至今", "现在") else _norm_month(end)
        # school / degree / major
        for line in block.splitlines():
            line = line.strip()
            if not line:
                continue
            if not item["school"] and ("大学" in line or "学院" in line):
                for tok in re.split(r"[\s|·,，、]+", line):
                    if "大学" in tok or "学院" in tok:
                        item["school"] = tok[:20]
                        break
                if not item["school"]:
                    item["school"] = line[:20]
            for deg in ("博士", "硕士", "本科", "大专", "高中"):
                if deg in line and not item["degree"]:
                    item["degree"] = deg
            if "专业" in line and not item["major"]:
                mm = re.search(r"([\u4e00-\u9fa5A-Za-z]+)\s*专业", line)
                if mm:
                    item["major"] = mm.group(1)
        if item["school"] or item["degree"]:
            items.append(item)
    return items


def _parse_experience_block(chunk: str) -> list[dict]:
    items: list[dict] = []
    blocks = [b.strip() for b in re.split(r"\n\s*\n", chunk) if b.strip()]
    for block in blocks[:6]:
        item = {"company": "", "position": "",
                "period_start": "", "period_end": "", "desc": ""}
        m = DATE_RANGE_RE.search(block)
        if m:
            item["period_start"] = _norm_month(m.group(1))
            end = m.group(2)
            item["period_end"] = "" if end in ("至今", "现在") else _norm_month(end)
        lines = [l.strip() for l in block.splitlines() if l.strip()]
        # first non-date line often has "公司 职位"
        head = next((l for l in lines if not DATE_RANGE_RE.search(l)), "")
        if head:
            parts = re.split(r"[\s|·,，\-／/]+", head, maxsplit=1)
            item["company"] = parts[0][:30]
            if len(parts) > 1:
                item["position"] = parts[1][:30]
        # rest is desc
        rest = "\n".join(l for l in lines if l != head and not DATE_RANGE_RE.fullmatch(l))
        item["desc"] = rest[:500]
        if item["company"] or item["desc"]:
            items.append(item)
    return items


def _parse_project_block(chunk: str) -> list[dict]:
    items: list[dict] = []
    blocks = [b.strip() for b in re.split(r"\n\s*\n", chunk) if b.strip()]
    for block in blocks[:6]:
        item = {"name": "", "role": "", "period_start": "",
                "period_end": "", "link": "", "desc": ""}
        m = DATE_RANGE_RE.search(block)
        if m:
            item["period_start"] = _norm_month(m.group(1))
            end = m.group(2)
            item["period_end"] = "" if end in ("至今", "现在") else _norm_month(end)
        lines = [l.strip() for l in block.splitlines() if l.strip()]
        head = next((l for l in lines if not DATE_RANGE_RE.search(l)), "")
        if head:
            item["name"] = head[:40]
        rest = "\n".join(l for l in lines if l != head and not DATE_RANGE_RE.fullmatch(l))
        item["desc"] = rest[:500]
        if item["name"] or item["desc"]:
            items.append(item)
    return items


def _parse_simple_block(chunk: str) -> list[dict]:
    """For competitions / certificates — one per line or paragraph."""
    items: list[dict] = []
    for line in re.split(r"[\n；;]", chunk):
        line = line.strip(" ·-•*\t")
        if len(line) < 2 or len(line) > 80:
            continue
        items.append({"name": line, "desc": ""})
        if len(items) >= 8:
            break
    return items


NAME_STOPWORDS = {
    "简历", "个人简历", "求职简历", "中文简历", "英文简历", "应聘",
    "姓名", "性别", "年龄", "民族", "籍贯", "政治", "政治面貌",
    "电话", "手机", "手机号", "电子邮箱", "邮箱", "邮件", "邮编",
    "地址", "住址", "家庭", "家庭住址", "联系方式", "出生", "生日",
    "学历", "专业", "学校", "院校", "学院", "本科", "硕士", "博士", "大专",
    "教育", "教育背景", "教育经历", "工作", "工作经历", "工作经验",
    "实习", "实习经历", "实习经验", "项目", "项目经历", "项目经验",
    "技能", "专业技能", "技能特长", "特长", "证书", "竞赛", "获奖",
    "语言", "语言能力", "自我评价", "自我介绍", "求职意向", "意向",
    "个人信息", "基本信息", "联系", "联系电话", "英语", "计算机",
    "男", "女", "备注",
}

def _guess_name(text: str) -> str:
    """Try multiple strategies to extract a Chinese or Western name."""
    head = text[:1000]
    # 1) Explicit label
    m = re.search(r"姓\s*名\s*[:：]?\s*([\u4e00-\u9fa5]{2,5}|[A-Za-z][A-Za-z .]{1,30})", head)
    if m:
        cand = m.group(1).strip()
        if cand not in NAME_STOPWORDS:
            return cand
    # 2) First non-empty line: if the whole line (stripped) is 2-5 Chinese chars
    #    and not a stopword, it's almost certainly the name (big title at top)
    lines = [l.strip() for l in head.splitlines() if l.strip()]
    if lines:
        first = re.sub(r"\s+", "", lines[0])
        if re.fullmatch(r"[\u4e00-\u9fa5]{2,5}", first) and first not in NAME_STOPWORDS:
            return first

    # 3) Scan the first ~30 non-empty lines; look at each whitespace-separated
    #    token and pick the first one that looks like a name
    for line in lines[:30]:
        # Split on whitespace, tabs, pipes, or common separators
        tokens = re.split(r"[\s|·•／/、,，]+", line)
        for tok in tokens:
            tok = tok.strip(" :：()（）[]【】")
            if not tok:
                continue
            # Chinese name: 2-4 Han chars, not a stopword
            if re.fullmatch(r"[\u4e00-\u9fa5]{2,4}", tok) and tok not in NAME_STOPWORDS:
                # Skip obvious non-names: must not be all same char ("男男"),
                # and shouldn't follow a label like "专业" etc
                return tok
            # Western name: two capitalized Latin tokens
        # Also try two-token Latin name on this line
        mm = re.fullmatch(r"([A-Z][a-z]+)\s+([A-Z][a-z]+)", line)
        if mm:
            return f"{mm.group(1)} {mm.group(2)}"
    return ""


def _heuristic_parse(text: str) -> dict:
    """Regex/keyword fallback so the form is populated without any LLM."""
    out: dict = {
        "basic": {}, "education": [], "internships": [],
        "work_experiences": [], "projects": [], "competitions": [],
        "certificates": [], "skills": [], "self_eval": "",
    }
    basic = out["basic"]
    m = EMAIL_RE.search(text)
    if m:
        basic["email"] = m.group(0)
    m = PHONE_RE.search(text)
    if m:
        basic["phone"] = m.group(0)
    name = _guess_name(text)
    if name:
        basic["name"] = name
    # Gender
    head_text = text[:500]
    if re.search(r"性\s*别[:：]?\s*男", head_text) or " 男 " in head_text:
        basic["gender"] = "男"
    elif re.search(r"性\s*别[:：]?\s*女", head_text) or " 女 " in head_text:
        basic["gender"] = "女"

    sections = _split_sections(text)

    if sections.get("education"):
        out["education"] = _parse_education_block(sections["education"])
    if not out["education"]:
        # fallback: pick first line containing 大学/学院
        for line in text.splitlines():
            if "大学" in line or "学院" in line:
                out["education"] = [{
                    "school": line.strip()[:30],
                    "degree": "本科" if "本科" in text else "",
                    "major": "", "college": "", "period_start": "",
                    "period_end": "", "advisor": "", "direction": "",
                }]
                break

    if sections.get("internships"):
        out["internships"] = _parse_experience_block(sections["internships"])
    if sections.get("work_experiences"):
        out["work_experiences"] = _parse_experience_block(sections["work_experiences"])
    if sections.get("projects"):
        out["projects"] = _parse_project_block(sections["projects"])
    if sections.get("competitions"):
        out["competitions"] = _parse_simple_block(sections["competitions"])
    if sections.get("certificates"):
        out["certificates"] = _parse_simple_block(sections["certificates"])

    # Skills - keyword scan over full text, plus anything in the skills section
    lower = text.lower()
    skills = [s for s in RESUME_SKILL_LEXICON if s.lower() in lower]
    out["skills"] = skills[:20]

    # Self eval
    if sections.get("self_eval"):
        out["self_eval"] = sections["self_eval"][:500]

    return out


def _merge_parsed(parsed: dict, user: str = "demo") -> None:
    """Only overwrite fields that the parser actually filled."""
    if not isinstance(parsed, dict):
        return
    cur = store.get(user).get("profile", {})
    if parsed.get("basic"):
        new_basic = {
            **(cur.get("basic") or {}),
            **{k: v for k, v in parsed["basic"].items() if v},
        }
        store.patch("profile.basic", new_basic, user=user)
    for key in ("education", "internships", "work_experiences",
                "projects", "experiences", "works", "competitions",
                "certificates", "languages", "socials", "skills"):
        if parsed.get(key):
            store.patch(f"profile.{key}", parsed[key], user=user)
    if parsed.get("self_eval"):
        store.patch("profile.self_eval", parsed["self_eval"], user=user)
    if parsed.get("soft_skills"):
        store.patch("profile.soft_skills", parsed["soft_skills"], user=user)


@app.post("/api/resume/upload")
async def resume_upload(request: Request, file: UploadFile = File(...)):
    content = (await file.read())[:500000]
    text = _extract_text(content, file.filename or "")
    print(f"[resume] {file.filename} bytes={len(content)} text_len={len(text)}")
    print(f"[resume] text head: {text[:200]!r}")
    # Try LLM first; fall back to heuristic if empty / offline stub
    prompt = (
        "请将以下简历文本解析为 JSON，字段包括："
        "basic{name,gender,phone,email,id_no,expected_cities},"
        "education[{period_start,period_end,school,degree,college,major,direction,advisor}],"
        "internships[{company,position,period_start,period_end,desc}],"
        "work_experiences[{company,position,period_start,period_end,desc}],"
        "projects[{name,role,period_start,period_end,link,desc}],"
        "competitions[{name,desc}],certificates[{name,desc}],"
        "languages[{language,proficiency}],self_eval,"
        "socials[{platform,url}],skills[]。只输出 JSON。\n\n"
        f"{text[:6000]}"
    )
    # Always run the deterministic heuristic first so we have a baseline
    heur = _heuristic_parse(text)
    # Only call LLM if a real API key is configured — the offline stub would
    # otherwise return misleading default values that overwrite heuristics.
    if llm.API_KEY:
        parsed = llm.chat_json("你是严谨的简历结构化解析器，只输出 JSON。", prompt)
        if not isinstance(parsed, dict):
            parsed = {}
    else:
        parsed = {}
    # Merge: heuristic fills anything the LLM left empty
    for k, v in heur.items():
        if not v:
            continue
        cur_v = parsed.get(k)
        if isinstance(v, dict):
            if not isinstance(cur_v, dict):
                parsed[k] = dict(v)
            else:
                for kk, vv in v.items():
                    if vv and not cur_v.get(kk):
                        cur_v[kk] = vv
        elif isinstance(v, list):
            if not cur_v:
                parsed[k] = v
        elif isinstance(v, str):
            if not cur_v:
                parsed[k] = v
    print(f"[resume] parsed basic={parsed.get('basic')} skills={parsed.get('skills')}")
    u = _current_user(request)
    _merge_parsed(parsed, user=u)
    store.patch("resume", {"filename": file.filename, "parsed": parsed}, user=u)
    return {"parsed": parsed, "profile": store.get(u).get("profile", {})}


class ProfileIn(BaseModel):
    basic: dict | None = None
    education: list | None = None
    internships: list | None = None
    work_experiences: list | None = None
    projects: list | None = None
    experiences: list | None = None  # legacy
    works: list | None = None
    competitions: list | None = None
    certificates: list | None = None
    languages: list | None = None
    self_eval: str | None = None
    socials: list | None = None
    skills: list | None = None
    soft_skills: dict | None = None
    intent: dict | None = None


@app.get("/api/profile")
def get_profile(request: Request):
    u = _current_user(request)
    s = store.get(u)
    prof = s.get("profile", {})
    completeness = _completeness(prof)
    competitiveness = _competitiveness(prof)
    return {
        "profile": prof,
        "completeness": completeness,
        "competitiveness": competitiveness,
        "progress": progress(s),
    }


@app.put("/api/profile")
def put_profile(body: ProfileIn, request: Request):
    u = _current_user(request)
    for k, v in body.model_dump(exclude_none=True).items():
        store.patch(f"profile.{k}", v, user=u)
    return get_profile(request)


def _completeness(p: dict) -> int:
    total, got = 0, 0
    # basic (name required for a tick)
    total += 1
    if (p.get("basic") or {}).get("name"):
        got += 1
    for k in ("education", "skills", "soft_skills", "intent"):
        total += 1
        v = p.get(k)
        if v and (isinstance(v, dict) and v or isinstance(v, list) and v):
            got += 1
    # "任一经历" 视为一项
    total += 1
    if any((p.get(k) or []) for k in
           ("internships", "work_experiences", "projects", "experiences")):
        got += 1
    # self_eval
    total += 1
    if (p.get("self_eval") or "").strip():
        got += 1
    return int(got / max(1, total) * 100)


def _competitiveness(p: dict) -> dict:
    # 档案几乎全空时直接返回 0，避免默认显示 40 误导用户
    has_any = bool(
        (p.get("skills") or [])
        or (p.get("education") or [])
        or (p.get("internships") or [])
        or (p.get("work_experiences") or [])
        or (p.get("projects") or [])
        or ((p.get("basic") or {}).get("name"))
        or (p.get("intent") or {})
        or (p.get("self_eval") or "").strip()
    )
    if not has_any:
        return {"score": 0, "comment": "请先完善档案，系统才能评估你的就业竞争力"}

    skills = p.get("skills") or []
    soft = p.get("soft_skills") or {}
    base = 40
    if soft:
        base += int(sum(soft.values()) / max(1, len(soft)) * 0.2)
    score = max(0, min(100, base + min(40, len(skills) * 3)))
    tag = "在互联网行业颇具潜力" if score > 70 else "基础扎实，建议继续补充项目经验"
    return {"score": score, "comment": tag}


# ---------- assessments ----------
class AssessIn(BaseModel):
    type: str  # holland | multi_intel | values | mbti
    result: Any  # list for holland/mi/values, list of floats for mbti


@app.post("/api/assessment")
def set_assess(body: AssessIn, request: Request):
    if body.type not in ("holland", "multi_intel", "values", "mbti"):
        raise HTTPException(400, "bad type")
    u = _current_user(request)
    store.patch(f"assessments.{body.type}", body.result, user=u)
    return {"ok": True, "progress": progress(store.get(u))}


@app.get("/api/assessment")
def get_assess(request: Request):
    return store.get(_current_user(request)).get("assessments", {})


HOLLAND_NAMES = ["R 实际型", "I 研究型", "A 艺术型", "S 社会型", "E 企业型", "C 常规型"]
VALUE_NAMES = [
    "成就感", "经济回报", "稳定安全", "自主独立", "创造创新",
    "社会影响", "人际关系", "学习成长", "生活平衡", "声望地位",
]
MI_NAMES = ["语言", "逻辑-数学", "空间", "音乐", "身体-运动", "人际", "内省", "自然观察"]
MBTI_DESCRIPTIONS = {
    "INTJ": "战略架构师 · 独立深度工作，系统设计，长期规划",
    "INTP": "逻辑探索者 · 抽象研究，技术钻研，理论分析",
    "ENTJ": "目标指挥官 · 领导团队，战略决策，高压环境",
    "ENTP": "创新辩论者 · 产品创新，商业模式探索，跨界整合",
    "INFJ": "洞察引导者 · 用户研究，心理支持，内容创作",
    "INFP": "理想探求者 · 创意写作，UX 设计，公益方向",
    "ENFJ": "人际协调者 · 团队管理，用户运营，教育培训",
    "ENFP": "热情连接者 · 市场营销，产品运营，社群运营",
    "ISTJ": "稳健执行者 · 数据分析，质量管理，项目控制",
    "ISFJ": "细心守护者 · 客户成功，测试，行政支持",
    "ESTJ": "秩序管理者 · 项目管理，运营管理，流程优化",
    "ESFJ": "和谐维护者 · HR，客服，团队协调",
    "ISTP": "务实技匠 · 后端开发，运维，嵌入式开发",
    "ISFP": "灵活创作者 · 视觉设计，前端开发，内容制作",
    "ESTP": "现实行动者 · 销售，增长运营，商务拓展",
    "ESFP": "活力表演者 · 社区运营，直播，品牌推广",
}


def _topk(names: list[str], scores: list, k: int = 3) -> list[tuple[str, int]]:
    pairs = [(names[i], int(scores[i] or 0)) for i in range(min(len(names), len(scores)))]
    pairs.sort(key=lambda x: x[1], reverse=True)
    return pairs[:k]


def _mbti_from_scores(s: list) -> str:
    if not s or len(s) < 4:
        return ""
    letters = [
        "E" if s[0] < 0.5 else "I",
        "S" if s[1] < 0.5 else "N",
        "T" if s[2] < 0.5 else "F",
        "J" if s[3] < 0.5 else "P",
    ]
    return "".join(letters)


def _holland_code(s: list) -> str:
    codes = ["R", "I", "A", "S", "E", "C"]
    if not s or len(s) < 6:
        return ""
    ranked = sorted(range(6), key=lambda i: (-s[i], i))
    return "".join(codes[i] for i in ranked[:3])


@app.get("/api/assessment/interpret")
def interpret_assess(request: Request):
    a = store.get(_current_user(request)).get("assessments", {})
    holland = a.get("holland") or []
    mi = a.get("multi_intel") or []
    values = a.get("values") or []
    mbti = a.get("mbti") or []

    if not any([holland, mi, values, mbti]):
        raise HTTPException(400, "暂无测评结果，请先完成至少一项测评")

    # Deterministic summary (always computed so we have a baseline)
    holland_top = _topk(HOLLAND_NAMES, holland) if holland else []
    values_top = _topk(VALUE_NAMES, values) if values else []
    mi_top = _topk(MI_NAMES, mi) if mi else []
    code = _holland_code(holland)
    mtype = _mbti_from_scores(mbti)

    # Build a descriptive prompt for the LLM
    lines: list[str] = []
    if holland:
        lines.append(
            "霍兰德代码：" + (code or "?") + "；前三兴趣："
            + "、".join(f"{n}({s}分)" for n, s in holland_top)
        )
    if mi:
        lines.append(
            "多元智能 Top 3：" + "、".join(f"{n}({s}分)" for n, s in mi_top)
        )
    if values:
        lines.append(
            "核心价值观 Top 3：" + "、".join(f"{n}({s}分)" for n, s in values_top)
        )
    if mtype:
        lines.append("MBTI：" + mtype + "（" + MBTI_DESCRIPTIONS.get(mtype, "") + "）")

    base_summary = "\n".join(lines)

    prompt = (
        "你是一位融合心理学和职业规划专业背景的生涯顾问。基于以下测评结果，"
        "为该学生生成一份 250-400 字的自我认识解析，需包含：\n"
        "1) 综合画像关键词一句话（如：逻辑深度+人际温度=产品技术复合型）\n"
        "2) 核心兴趣方向与对应职业领域\n"
        "3) 驱动力和价值观倾向\n"
        "4) 认知优势组合\n"
        "5) 工作风格标签\n"
        "6) 可能存在的冲突或盲区提醒\n"
        "语言亲切自然，避免术语堆砌，不直接推荐具体岗位。\n\n"
        f"测评摘要：\n{base_summary}"
    )

    narrative = llm.chat(
        "你是资深的职业规划顾问，擅长用温暖而精准的语言帮助学生理解自己。",
        prompt,
    )

    # If llm returned an offline stub, craft a deterministic narrative instead
    if not narrative or "offline-stub" in narrative:
        parts: list[str] = []

        # ── 1. 综合画像一句话 ──
        profile_tags: list[str] = []
        if code:
            h_letter = code[0]
            h_short = {"R": "动手实操", "I": "深度研究", "A": "创意表达", "S": "人际支持", "E": "领导驱动", "C": "有序执行"}.get(h_letter, "")
            profile_tags.append(h_short)
        if mi_top:
            profile_tags.append(f"{mi_top[0][0]}优势")
        if mtype:
            ei = "内敛专注" if mtype[0] == "I" else "外向行动"
            profile_tags.append(ei)
        if profile_tags:
            parts.append(f"**【综合画像】** {'＋'.join(profile_tags)}型人才。")

        # ── 2. 兴趣方向深度解读（霍兰德） ──
        if code and holland_top:
            lead_name = holland_top[0][0]
            lead_letter = code[0]
            second_letter = code[1] if len(code) > 1 else ""

            # 霍兰德类型的深度特征
            h_deep = {
                "R": ("你倾向于通过具体操作和实践来解决问题，喜欢看到实实在在的成果。", "工程技术、运维、硬件开发、数据工程"),
                "I": ("你天生对'为什么'感兴趣，享受独立思考和深度分析的过程，更看重理解本质而非执行流程。", "数据分析、算法研究、心理学研究、咨询分析"),
                "A": ("你需要在工作中获得创造性的表达空间，重复性高的任务会让你感到窒息。", "内容创作、UI/UX 设计、品牌策划、视觉传达"),
                "S": ("你的职业满足感来自'帮到了别人'，人际连接是你的核心能量来源。", "教育培训、用户研究、HR、心理咨询、社区运营"),
                "E": ("你天然追求影响力和资源调动权，在有明确目标和竞争的环境里最能发挥。", "产品管理、商务拓展、项目管理、创业"),
                "C": ("你在有序、可预期的环境中表现最佳，擅长把复杂流程变成高效系统。", "财务分析、审计、数据处理、质量管理"),
            }
            desc, fields = h_deep.get(lead_letter, ("", ""))
            parts.append(
                f"**【兴趣方向】** 你的霍兰德代码是 **{code}**，主导类型 {lead_name}。{desc}"
                f"适配的职业领域包括：{fields}。"
            )

            # 双码组合解读
            combo_insight = {
                "IR": "你是少见的'研究+实操'双核型，既能做理论也能落地实现，技术研发类岗位最能发挥这种组合优势。",
                "IA": "研究型+艺术型的组合意味着你追求'有深度的创造'，适合用户体验研究、内容策略等需要洞察力+创意的方向。",
                "IS": "研究型+社会型组合在心理咨询、教育研究、用户研究领域有天然优势，你既能理解人，又能系统分析。",
                "IE": "研究型+企业型是典型的'技术管理者'画像——先靠专业能力立足，再向管理方向发展。",
                "IC": "研究型+常规型组合适合需要严谨分析的领域：数据分析、审计、金融风控、质量检测。",
                "AS": "艺术型+社会型意味着你希望工作既有创造性又有意义感，内容教育、公益设计、社群运营可能是你的甜区。",
                "AE": "艺术型+企业型是品牌营销、创意总监的典型画像——既有创意又有商业嗅觉。",
                "SE": "社会型+企业型组合适合需要'影响人'的岗位：培训管理、用户增长、团队领导。",
                "SC": "社会型+常规型组合适合服务导向+流程化的工作：HR、客户成功、行政管理。",
                "EC": "企业型+常规型是项目管理、运营管理的标准画像——目标导向+执行有序。",
                "RI": "实际型+研究型是工程师的经典组合，适合需要'动脑+动手'的技术岗位。",
                "RA": "实际型+艺术型组合少见但有价值，适合工业设计、前端开发、3D 建模等方向。",
            }
            pair = lead_letter + second_letter
            pair_rev = second_letter + lead_letter
            combo = combo_insight.get(pair) or combo_insight.get(pair_rev, "")
            if combo:
                parts.append(combo)

        # ── 3. 驱动力深度解读（价值观 × MBTI） ──
        if values_top:
            v_names = [n for n, _ in values_top]
            parts.append(
                f"**【核心驱动力】** 你最看重 **{'、'.join(v_names)}**。"
            )
            # 价值观组合洞察
            v_set = set(v_names)
            v_insights: list[str] = []
            if "成就感" in v_set and "学习成长" in v_set:
                v_insights.append("成就感+学习成长说明你是内驱型人格，不需要外部压力就能持续进步，但要注意避免完美主义带来的自我消耗。")
            if "经济回报" in v_set and "稳定安全" in v_set:
                v_insights.append("经济回报+稳定安全说明你是务实型选择者，建议优先看大厂或成熟企业的成熟岗位，不适合早期创业。")
            if "自主独立" in v_set and "创造创新" in v_set:
                v_insights.append("自主独立+创造创新意味着你需要高自由度的工作环境，传统层级分明的组织会让你感到压抑。")
            if "社会影响" in v_set and "人际关系" in v_set:
                v_insights.append("社会影响+人际关系驱动说明你从'对他人有价值'中获取能量，教育、公益、咨询类方向最契合这种驱动力。")
            if "学习成长" in v_set and "创造创新" in v_set:
                v_insights.append("学习成长+创造创新组合适合快速迭代的行业（互联网、AI、新媒体），在变化中你反而如鱼得水。")
            if "生活平衡" in v_set:
                v_insights.append("你把生活平衡列为核心价值观，这不是'没有野心'，而是清楚自己的边界。选岗位时建议把加班文化作为重要筛选项。")

            # MBTI × 价值观交叉
            if mtype:
                if mtype[2] == "F" and "经济回报" in v_set:
                    v_insights.append(f"你的 MBTI 是情感型（{mtype}），但价值观又重视经济回报——这说明你希望'做有温度的事，同时过有品质的生活'，而非纯粹追求高薪。")
                if mtype[2] == "T" and "人际关系" in v_set:
                    v_insights.append(f"你的 MBTI 是思考型（{mtype}），但价值观看重人际关系——你可能在团队中扮演'理性的关怀者'角色，用逻辑帮人解决问题而非单纯共情。")
                if mtype[3] == "J" and "自主独立" in v_set:
                    v_insights.append(f"你是 J 型（喜欢计划和结构），但追求自主独立——你需要的不是混乱的自由，而是'在自己设定的框架里自主行动'。")

            if v_insights:
                parts.append("\n".join(v_insights[:2]))

        # ── 4. 认知优势组合解读（多元智能 × 霍兰德） ──
        if mi_top:
            mi_names = [n for n, _ in mi_top]
            parts.append(
                f"**【认知优势】** 你的优势智能是 **{'、'.join(mi_names)}**。"
            )
            mi_set = set(mi_names)
            mi_insights: list[str] = []

            # 智能组合的协同效应
            if "语言" in mi_set and "人际" in mi_set:
                mi_insights.append("语言+人际是'沟通型人才'的标配——你不仅理解人，还能把理解清晰地表达出来，这在产品、运营、咨询领域是稀缺能力。")
            if "逻辑-数学" in mi_set and "语言" in mi_set:
                mi_insights.append("逻辑+语言的组合让你既能分析问题又能讲清楚分析结果，非常适合数据分析师、咨询顾问、技术写作等'桥梁型'岗位。")
            if "空间" in mi_set and "逻辑-数学" in mi_set:
                mi_insights.append("空间+逻辑组合意味着你擅长把抽象概念可视化，前端开发、系统架构、数据可视化是你的优势战场。")
            if "内省" in mi_set and "人际" in mi_set:
                mi_insights.append("内省+人际说明你既懂自己又懂别人，这种'双向共情'在心理咨询、用户研究、团队管理中极有价值。")
            if "身体-运动" in mi_set and "空间" in mi_set:
                mi_insights.append("身体-运动+空间组合适合需要手脑协调的方向：工业设计、硬件开发、实验操作。")
            if "音乐" in mi_set and "语言" in mi_set:
                mi_insights.append("音乐+语言组合说明你对节奏和表达都敏感，播客、短视频创作、声音设计可能是你的隐藏赛道。")

            # 智能 × 霍兰德的协同/冲突
            if code:
                lead_letter = code[0]
                if lead_letter == "I" and "人际" in mi_set:
                    mi_insights.append("有趣的是，你的主导兴趣是研究型，但人际智能很突出——这说明你不是'只会埋头做研究的书呆子'，而是能把研究成果转化为可理解的洞察，用户研究、数据咨询方向会很适合你。")
                if lead_letter == "S" and "逻辑-数学" in mi_set:
                    mi_insights.append("你的主导兴趣是社会型，但逻辑智能突出——你是那种'用数据说话的暖心人'，HR 数据分析、教育测评设计可能是你的差异化方向。")
                if lead_letter == "A" and "逻辑-数学" in mi_set:
                    mi_insights.append("艺术型兴趣+逻辑智能是'理性创意者'——你的创意不是天马行空，而是有数据支撑的，增长运营、A/B 测试驱动的设计方向会让你的两种能力同时发光。")

            if mi_insights:
                parts.append("\n".join(mi_insights[:2]))

        # ── 5. 工作风格解读（MBTI 深度） ──
        if mtype:
            mbti_desc = MBTI_DESCRIPTIONS.get(mtype, "")
            parts.append(
                f"**【工作风格】** 你的 MBTI 是 **{mtype}**（{mbti_desc.split('·')[0].strip()}）。"
            )
            # MBTI 深度解读
            style_notes: list[str] = []
            ei = "你更倾向于独立思考和深度工作，开放式办公和频繁会议会消耗你的能量，建议寻找允许远程或有独立空间的工作环境。" if mtype[0] == "I" else "你在团队互动中获取能量，长期独自工作会让你感到孤立，选择团队协作密集的岗位会更有活力。"
            style_notes.append(ei)

            sn = "你更关注具体事实和当下任务，擅长执行和落地，但在面试中要注意展现你对趋势和战略的思考。" if mtype[1] == "S" else "你更关注可能性和未来趋势，适合需要前瞻思考的岗位，但要注意避免'想太多做太少'。"
            style_notes.append(sn)

            tf = "你决策时优先考虑逻辑和效率，在技术和分析岗位上这是优势，但在团队管理中要有意识地关注他人感受。" if mtype[2] == "T" else "你决策时会考虑人的感受和价值观，这在用户导向的岗位中是天然优势，但要注意不因共情而影响客观判断。"
            style_notes.append(tf)

            parts.append(style_notes[0] + " " + style_notes[1] + " " + style_notes[2])

        # ── 6. 潜在盲区 ──
        blindspots: list[str] = []
        if code and mtype:
            if code[0] == "I" and mtype[0] == "I":
                blindspots.append("你的兴趣和性格都偏向内向独立，要警惕'社交肌肉萎缩'——定期主动参加行业活动或社群，保持人脉的活跃度。")
            if code[0] == "E" and mtype[0] == "I":
                blindspots.append("你的兴趣偏企业型（追求影响力），但性格偏内向——你可能更适合'幕后影响者'的角色，如战略顾问、产品负责人，而非一线销售。")
            if code[0] == "A" and mtype[3] == "J":
                blindspots.append("你的兴趣偏艺术创意型，但 J 型性格喜欢计划和结构——找一个'有框架的创意岗位'（如品牌视觉规范设计）会比纯自由创作更适合你。")
        if values_top:
            v_set = set(n for n, _ in values_top)
            if "经济回报" in v_set and "生活平衡" in v_set:
                blindspots.append("你同时追求高收入和生活平衡，在现实中这两者存在张力——建议明确底线：多少收入以上你愿意牺牲多少生活时间？")
            if "自主独立" in v_set and "稳定安全" in v_set:
                blindspots.append("你既想自主又想稳定——纯体制内可能太僵化，纯自由职业又太不确定，大厂内部创新岗或成熟公司的独立项目组可能是最佳平衡点。")

        if blindspots:
            parts.append("**【潜在盲区】**\n" + "\n".join(f"- {b}" for b in blindspots[:2]))

        if not parts:
            parts.append("测评数据不足，请先完成更多测评以获得更准确的解析。")
        parts.append(
            "**提醒**：以上是基于四项测评交叉分析的参考画像。"
            "测评揭示的是你的倾向和潜力，而非命运——最终的职业决策还需要结合你的真实经历、市场机会与个人阶段综合判断。"
        )
        narrative = "\n\n".join(parts)

    return {
        "narrative": narrative,
        "summary": {
            "holland_code": code,
            "holland_top": holland_top,
            "values_top": values_top,
            "mi_top": mi_top,
            "mbti": mtype,
            "mbti_label": MBTI_DESCRIPTIONS.get(mtype, ""),
        },
    }


# ---------- file storage ----------
import uuid as _uuid

UPLOADS_DIR = ROOT / "app" / "data" / "uploads"
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

# 允许的文件扩展名和最大大小（300 MB）
ALLOWED_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".ppt", ".pptx", ".xls", ".xlsx",
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg",
    ".zip", ".rar", ".7z", ".tar", ".gz",
    ".mp4", ".mp3", ".wav", ".avi", ".mov",
    ".txt", ".md", ".csv", ".json",
}
MAX_FILE_SIZE = 300 * 1024 * 1024  # 300 MB


@app.post("/api/files/upload")
async def upload_file(file: UploadFile = File(...)):
    """上传文件到本地 uploads 目录，返回文件 ID 和下载 URL。"""
    if not file.filename:
        raise HTTPException(400, "文件名为空")
    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, f"不支持的文件类型：{ext}")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(400, f"文件大小超过限制（最大 {MAX_FILE_SIZE // 1024 // 1024} MB）")

    # 用 UUID 做唯一文件名，保留原始扩展名
    file_id = _uuid.uuid4().hex[:12]
    safe_name = f"{file_id}{ext}"
    (UPLOADS_DIR / safe_name).write_bytes(content)

    return {
        "file_id": file_id,
        "filename": file.filename,
        "size": len(content),
        "url": f"/api/files/{safe_name}",
    }


from fastapi.responses import FileResponse as _FileResponse


@app.get("/api/files/{filename}")
def serve_file(filename: str):
    """下载/预览已上传的文件。"""
    # 安全检查：防止路径穿越
    safe = Path(filename).name
    fp = UPLOADS_DIR / safe
    if not fp.exists() or not fp.is_file():
        raise HTTPException(404, "文件不存在")
    return _FileResponse(fp, filename=safe)


# ---------- deep self-awareness (思索工作的步骤 · 四步法) ----------

DEEP_FIELD_LABELS = {
    "good_at": "我擅长的",
    "like": "我喜欢的",
    "pursue": "我追求的",
    "focus": "聚焦方向",
    "target_jobs": "目标岗位",
    "conditions": "过滤条件",
}


class DeepAssessIn(BaseModel):
    answers: dict  # {good_at, like, pursue, focus, target_jobs, conditions}


class Step2In(BaseModel):
    good_at: str = ""
    like: str = ""
    pursue: str = ""


def _extract_keywords(text: str) -> list[str]:
    """从用户文本中提取关键词（按中文逗号/顿号/换行/空格拆分）"""
    tokens = re.split(r"[,，、；;\n\r\s/|]+", (text or "").strip())
    return [t.strip() for t in tokens if t.strip()]


# ---- 语义标签体系：用于从自然语言中识别用户的能力/兴趣/追求 ----

_TRAIT_TAGS: dict[str, list[str]] = {
    # 能力 / 技能类
    "写作": ["写作", "写文章", "文案", "文笔", "撰写", "编辑", "科普", "宣传文", "文字"],
    "阅读": ["阅读", "读书", "看书"],
    "分析": ["分析", "逻辑", "推理", "体系", "框架", "思维", "结构"],
    "审美": ["审美", "美感", "艺术", "美学", "美相关", "跟美"],
    "心理学": ["心理", "心理学", "心理活动", "心理咨询", "情绪"],
    "沟通": ["沟通", "表达", "交流", "协调", "说服", "演讲", "脱口秀", "口才"],
    "学习能力": ["学东西", "学习", "学新东西", "进步", "成长", "不断地学"],
    "检索信息": ["检索", "搜索", "信息", "查找", "调研"],
    "编程": ["编程", "代码", "开发", "程序", "软件", "Python", "Java", "前端", "后端"],
    "数据": ["数据", "数据分析", "统计", "量化"],
    "设计": ["设计", "创意", "UI", "UX", "视觉", "交互", "平面", "美术"],
    "管理": ["管理", "领导", "带团队", "组织"],
    "运营": ["运营", "推广", "增长", "内容运营"],
    "营销": ["营销", "市场", "销售", "品牌", "策划", "商务"],
    "教育": ["教育", "教学", "培训", "辅导", "授课", "老师"],
    "研究": ["研究", "科研", "学术", "实验", "论文"],
    "金融": ["金融", "财务", "会计", "投资", "银行", "审计", "经济", "理财"],
    "种植": ["种地", "种植", "农业", "种花", "园艺"],
    "自然": ["大自然", "自然", "户外", "生态", "环保"],
    "综艺": ["综艺", "娱乐", "看综艺", "节目"],
    "安静": ["安静", "独处", "内向", "低调"],
    "助人": ["帮助", "帮人", "输出", "服务", "公益", "志愿"],
}

# 方向 → 需要匹配的标签组合 + 对应的岗位推荐
_DIRECTION_RULES: list[dict] = [
    {
        "name": "内容创作 / 新媒体方向",
        "require_any": [("写作", "阅读"), ("写作", "审美"), ("写作", "沟通")],
        "boost_tags": ["学习能力", "检索信息"],
        "jobs": ["内容运营", "新媒体运营", "文案策划", "编辑", "自媒体"],
        "reason": "你擅长写作且喜欢阅读/审美，这类能力在内容创作领域非常有价值",
    },
    {
        "name": "心理咨询 / 教育方向",
        "require_any": [("心理学", "沟通"), ("心理学", "助人"), ("心理学", "分析")],
        "boost_tags": ["学习能力", "阅读"],
        "jobs": ["心理咨询师", "职业规划师", "培训师", "教育产品经理", "用户研究"],
        "reason": "你对心理学有兴趣且善于沟通/分析，适合需要理解人的岗位",
    },
    {
        "name": "用户研究 / 产品方向",
        "require_any": [("分析", "沟通"), ("心理学", "数据"), ("分析", "审美"), ("检索信息", "分析")],
        "boost_tags": ["写作", "学习能力"],
        "jobs": ["用户研究", "产品经理", "数据分析师", "市场调研", "用户体验设计师"],
        "reason": "你兼具分析思维和人际理解能力，适合连接用户需求与产品设计的角色",
    },
    {
        "name": "数据分析 / 商业分析方向",
        "require_any": [("数据", "分析"), ("分析", "金融"), ("编程", "分析")],
        "boost_tags": ["学习能力", "检索信息"],
        "jobs": ["数据分析师", "商业分析师", "BI工程师", "量化分析", "风控分析"],
        "reason": "你擅长分析和逻辑思考，数据驱动的岗位能发挥你的优势",
    },
    {
        "name": "技术开发方向",
        "require_any": [("编程", "学习能力"), ("编程", "数据"), ("编程", "分析")],
        "boost_tags": ["检索信息"],
        "jobs": ["软件工程师", "前端开发", "后端开发", "全栈工程师", "算法工程师"],
        "reason": "你有编程基础且持续学习，技术岗位能给你不断成长的空间",
    },
    {
        "name": "设计 / 创意方向",
        "require_any": [("设计", "审美"), ("审美", "沟通"), ("设计", "学习能力")],
        "boost_tags": ["写作"],
        "jobs": ["UI设计师", "平面设计师", "交互设计师", "品牌设计", "视觉设计"],
        "reason": "你具备审美能力和创意思维，设计领域可以将这些天赋转化为价值",
    },
    {
        "name": "运营 / 营销方向",
        "require_any": [("运营", "写作"), ("营销", "沟通"), ("运营", "数据"), ("营销", "分析")],
        "boost_tags": ["审美", "学习能力"],
        "jobs": ["运营经理", "品牌策划", "活动运营", "社群运营", "市场推广"],
        "reason": "你善于沟通且有策划能力，运营/营销岗位需要这样的综合素质",
    },
    {
        "name": "教育培训方向",
        "require_any": [("教育", "沟通"), ("教育", "写作"), ("助人", "沟通"), ("助人", "学习能力")],
        "boost_tags": ["心理学", "分析"],
        "jobs": ["培训讲师", "课程设计", "教育咨询", "学科老师", "教育产品经理"],
        "reason": "你有帮助他人的热情，教育领域能让你的输出产生持久影响",
    },
    {
        "name": "研究 / 学术方向",
        "require_any": [("研究", "分析"), ("研究", "写作"), ("研究", "学习能力")],
        "boost_tags": ["阅读", "检索信息"],
        "jobs": ["研究员", "学术助理", "政策研究", "行业分析师", "智库研究员"],
        "reason": "你有扎实的研究和分析能力，适合需要深度思考的学术/研究岗位",
    },
    {
        "name": "金融 / 财务方向",
        "require_any": [("金融", "分析"), ("金融", "数据"), ("金融", "学习能力")],
        "boost_tags": ["检索信息"],
        "jobs": ["投资分析师", "风控专员", "财务分析", "基金经理助理", "审计"],
        "reason": "你对金融有兴趣且具备分析能力，金融行业可以满足你的成长和收入追求",
    },
]


def _scan_tags(text: str) -> set[str]:
    """从自然语言文本中扫描匹配的语义标签"""
    matched = set()
    lower = text.lower()
    for tag, keywords in _TRAIT_TAGS.items():
        for kw in keywords:
            if kw.lower() in lower:
                matched.add(tag)
                break
    return matched


def _offline_step2_analysis(good_at: str, like: str, pursue: str, assessments: dict) -> dict:
    """离线交叉分析：语义标签匹配 + 交叉推理 + 岗位推荐"""
    tags_good = _scan_tags(good_at)
    tags_like = _scan_tags(like)
    tags_pursue = _scan_tags(pursue)
    all_tags = tags_good | tags_like | tags_pursue

    # ---- 交叉分析 ----
    cross_good_like = tags_good & tags_like       # 擅长 ∩ 喜欢
    cross_like_pursue = tags_like & tags_pursue   # 喜欢 ∩ 追求
    cross_good_pursue = tags_good & tags_pursue   # 擅长 ∩ 追求
    cross_all = tags_good & tags_like & tags_pursue  # 三圆交集

    # ---- 测评数据 ----
    holland = assessments.get("holland")
    mbti = assessments.get("mbti")
    holland_top = []
    mbti_code = ""
    h_order = ["R", "I", "A", "S", "E", "C"]
    h_labels = {"R": "实际型", "I": "研究型", "A": "艺术型", "S": "社会型", "E": "企业型", "C": "常规型"}
    if holland:
        try:
            ranked = sorted(range(len(h_order)), key=lambda i: (holland[i] if i < len(holland) else 0), reverse=True)
            holland_top = [h_order[i] for i in ranked[:2] if i < len(h_order)]
        except Exception:
            pass
    if mbti:
        try:
            pairs = [("E", "I"), ("S", "N"), ("T", "F"), ("J", "P")]
            for i, (a, b) in enumerate(pairs):
                v = mbti[i] if i < len(mbti) else 0.5
                mbti_code += b if v >= 0.5 else a
        except Exception:
            pass

    # Holland 也可以增强方向判断
    holland_boost_tags: set[str] = set()
    for code in holland_top:
        if code == "I":
            holland_boost_tags |= {"分析", "研究", "学习能力"}
        elif code == "A":
            holland_boost_tags |= {"审美", "设计", "写作"}
        elif code == "S":
            holland_boost_tags |= {"沟通", "助人", "教育", "心理学"}
        elif code == "E":
            holland_boost_tags |= {"管理", "营销", "沟通"}
        elif code == "C":
            holland_boost_tags |= {"数据", "分析", "金融"}
        elif code == "R":
            holland_boost_tags |= {"编程", "种植", "自然"}

    # ---- 方向匹配打分 ----
    scored_directions: list[tuple[float, dict]] = []
    for rule in _DIRECTION_RULES:
        score = 0.0
        matched_pairs = []
        for pair in rule["require_any"]:
            a, b = pair
            # 核心：两个标签分别出现在不同维度才算交叉
            sources_a = []
            if a in tags_good: sources_a.append("擅长")
            if a in tags_like: sources_a.append("喜欢")
            if a in tags_pursue: sources_a.append("追求")
            sources_b = []
            if b in tags_good: sources_b.append("擅长")
            if b in tags_like: sources_b.append("喜欢")
            if b in tags_pursue: sources_b.append("追求")
            if sources_a and sources_b:
                score += 2.0
                matched_pairs.append((a, b, sources_a, sources_b))
            elif a in all_tags and b in all_tags:
                score += 1.0
                matched_pairs.append((a, b, sources_a or ["—"], sources_b or ["—"]))
            elif a in all_tags or b in all_tags:
                score += 0.3
        # boost
        for bt in rule.get("boost_tags", []):
            if bt in all_tags:
                score += 0.5
            if bt in holland_boost_tags:
                score += 0.3
        if score >= 1.0:
            scored_directions.append((score, {**rule, "_matched": matched_pairs}))

    scored_directions.sort(key=lambda x: -x[0])
    top_dirs = scored_directions[:3]

    # ---- 生成报告 ----
    parts: list[str] = []

    # 1) 能力画像
    parts.append("## 能力画像")
    if tags_good:
        parts.append(f"**你擅长的核心能力**：{'、'.join(tags_good)}")
    if tags_like:
        parts.append(f"**你真正喜欢的事**：{'、'.join(tags_like)}")
    if tags_pursue:
        parts.append(f"**你内心追求的**：{'、'.join(tags_pursue)}")
    parts.append("")

    # 2) 交叉点
    parts.append("## 交叉点分析")
    if cross_all:
        parts.append(f"**三维交集（擅长 ∩ 喜欢 ∩ 追求）**：**{'、'.join(cross_all)}**")
        parts.append("这是你最核心的方向——既擅长、又喜欢、还符合追求，优先围绕这些能力选择岗位。")
    if cross_good_like - cross_all:
        tags_gl = cross_good_like - cross_all
        parts.append(f"**擅长 ∩ 喜欢**：{'、'.join(tags_gl)} — 做起来既顺手又有乐趣，容易进入心流状态")
    if cross_like_pursue - cross_all:
        tags_lp = cross_like_pursue - cross_all
        parts.append(f"**喜欢 ∩ 追求**：{'、'.join(tags_lp)} — 有热情也有动力，但可能需要进一步提升技能")
    if cross_good_pursue - cross_all:
        tags_gp = cross_good_pursue - cross_all
        parts.append(f"**擅长 ∩ 追求**：{'、'.join(tags_gp)} — 有能力也有目标，但需要确认是否真的喜欢")
    if not cross_good_like and not cross_like_pursue and not cross_good_pursue:
        parts.append("三个维度目前没有直接的标签交叉，但这并不意味着没有方向。")
        parts.append("结合你描述的细节，系统从语义层面为你推理了以下方向。")
    parts.append("")

    # 3) 四维测评交叉分析
    values = assessments.get("values")
    mi = assessments.get("multi_intel")
    values_top = _topk(VALUE_NAMES, values)[:3] if values and len(values) >= 6 else []
    mi_top_list = _topk(MI_NAMES, mi)[:3] if mi and len(mi) >= 6 else []
    has_any_assess = holland_top or mbti_code or values_top or mi_top_list

    if has_any_assess:
        parts.append("## 测评交叉验证")

        # 霍兰德 × 头脑风暴 对齐
        if holland_top:
            h_names = [h_labels.get(c, c) for c in holland_top]
            h_tag_map = {"I": {"分析", "研究", "学习能力"}, "A": {"审美", "设计", "写作"}, "S": {"沟通", "助人", "教育", "心理学"}, "E": {"管理", "营销"}, "C": {"数据", "金融"}, "R": {"编程", "自然"}}
            aligned = h_tag_map.get(holland_top[0], set()) & all_tags
            if aligned:
                parts.append(f"你的霍兰德主导兴趣 **{h_names[0]}** 与你头脑风暴中提到的 **{'、'.join(aligned)}** 高度一致，说明你的自我认知和测评结果相互印证，方向可信度高。")
            else:
                parts.append(f"你的霍兰德主导兴趣 **{h_names[0]}**，但头脑风暴中没有出现直接对应的标签，可能存在'潜意识兴趣'尚未被你清晰表达——建议在后续探索中留意这个方向。")

        # 价值观 × 追求维度 对齐
        if values_top and tags_pursue:
            v_names = [n for n, _ in values_top]
            v_tag_map = {"成就感": {"学习能力", "分析"}, "经济回报": {"金融", "营销"}, "学习成长": {"学习能力", "研究"}, "创造创新": {"设计", "审美", "写作"}, "社会影响": {"助人", "教育"}, "人际关系": {"沟通", "助人"}, "自主独立": {"编程", "写作"}, "稳定安全": {"数据", "金融"}}
            v_aligned = set()
            for vn in v_names:
                v_aligned |= (v_tag_map.get(vn, set()) & tags_pursue)
            if v_aligned:
                parts.append(f"你的核心价值观 **{'、'.join(v_names)}** 与你「追求的」维度中 **{'、'.join(v_aligned)}** 相呼应——你知道自己想要什么，而且追求方式与内心驱动力一致。")
            else:
                parts.append(f"你的核心价值观是 **{'、'.join(v_names)}**，但「追求的」维度中没有直接体现——值得思考：你追求的目标是否真正匹配你的内在驱动力？")

        # 多元智能 × 擅长维度 对齐
        if mi_top_list and tags_good:
            mi_names = [n for n, _ in mi_top_list]
            mi_tag_map = {"语言": {"写作", "阅读", "沟通"}, "逻辑-数学": {"分析", "数据", "编程"}, "空间": {"设计", "审美"}, "人际": {"沟通", "助人", "教育"}, "内省": {"心理学", "安静"}, "音乐": {"审美"}, "自然观察": {"自然", "种植"}, "身体-运动": {"运营"}}
            mi_aligned = set()
            for mn in mi_names:
                mi_aligned |= (mi_tag_map.get(mn, set()) & tags_good)
            if mi_aligned:
                parts.append(f"你的优势智能 **{'、'.join(mi_names)}** 正好对应你擅长的 **{'、'.join(mi_aligned)}**——这些不只是你'会'的东西，更是你认知层面的天然优势，建议重点发展。")
            else:
                parts.append(f"你的优势智能是 **{'、'.join(mi_names)}**，这些认知优势可以作为你未来职业能力的底层支撑。")

        # MBTI × 整体风格
        if mbti_code:
            mbti_label = MBTI_DESCRIPTIONS.get(mbti_code, "")
            style_tag = mbti_label.split("·")[0].strip() if mbti_label else ""
            work_style = ""
            if mbti_code[0] == "I" and ("安静" in all_tags or "阅读" in all_tags):
                work_style = "你的 MBTI 偏内向，头脑风暴中也提到了安静/阅读等关键词，你在独立、专注的工作环境中会更高效。"
            elif mbti_code[0] == "E" and ("沟通" in all_tags or "助人" in all_tags):
                work_style = "你的 MBTI 偏外向，加上你擅长和喜欢沟通/助人，你在需要人际互动的岗位上会更有能量。"
            elif mbti_code[0] == "I" and "沟通" in all_tags:
                work_style = f"有意思的是，你的 MBTI 偏内向（{mbti_code}），但你却提到了沟通能力——你可能是'一对一深度沟通'型而非'社交达人'型，用户研究、咨询类岗位比大规模社交更适合你。"
            else:
                work_style = f"你的 MBTI 为 **{mbti_code}**（{style_tag}），这决定了你自然的工作节奏和协作方式。"
            parts.append(work_style)

        parts.append("")

    # 4) 推荐方向 + 岗位
    parts.append("## 推荐方向")
    all_recommended_jobs: list[str] = []
    if top_dirs:
        for i, (score, d) in enumerate(top_dirs, 1):
            stars = "★" * min(int(score), 5)
            parts.append(f"### {i}. {d['name']}  {stars}")
            parts.append(f"{d['reason']}。")
            if d.get("_matched"):
                evidence = []
                for a, b, sa, sb in d["_matched"][:2]:
                    evidence.append(f"「{a}」({'·'.join(sa)}) × 「{b}」({'·'.join(sb)})")
                parts.append(f"匹配依据：{' + '.join(evidence)}")
            parts.append(f"**推荐岗位**：{'、'.join(d['jobs'])}")
            all_recommended_jobs.extend(d["jobs"])
    else:
        # 没有匹配到规则，给通用建议
        parts.append("根据你的描述，暂时没有匹配到高置信度的方向。建议：")
        parts.append("- 把「我擅长的」写得更具体一些（技能、工具、曾做过的事）")
        parts.append("- 想想哪些事情让你忘记时间（这通常就是你真正喜欢的）")
        parts.append("- 追求可以从「5 年后想成为什么样的人」来思考")

    # 智慧库：注入与用户情况相关的原则（排除不适用的）
    skip_ids = {"action_first", "three_steps"}  # 有测评数据说明用户已在行动，不需要"先动起来"
    eligible = [p for p in _UNIVERSAL_PRINCIPLES if p.get("id") not in skip_ids]
    if eligible:
        chosen = random.sample(eligible, min(2, len(eligible)))
        parts.append("\n## 给你的建议")
        for p in chosen:
            parts.append(f"**{p['name']}**：{p['desc']}\n")

    directions = [d["name"] for _, d in top_dirs]
    focus_suggestion = "、".join(directions[:2]) if directions else ""
    return {
        "analysis": "\n".join(parts),
        "focus_suggestion": focus_suggestion,
        "directions": directions,
        "recommended_jobs": all_recommended_jobs[:10],
    }


def _offline_deep_analysis(answers: dict, assessments: dict | None = None) -> str:
    """离线生成四步法最终分析报告 —— 复用 step2 交叉分析引擎"""
    parts: list[str] = ["# 思索工作 · 职业方向分析报告\n"]

    good_at = (answers.get("good_at") or "").strip()
    like = (answers.get("like") or "").strip()
    pursue = (answers.get("pursue") or "").strip()
    focus = (answers.get("focus") or "").strip()
    target_jobs = answers.get("target_jobs") or []

    # ── 一、头脑风暴回顾 ──
    parts.append("## 一、头脑风暴\n")
    parts.append(f"**我擅长的**：{good_at or '（未填写）'}")
    parts.append(f"**我喜欢的**：{like or '（未填写）'}")
    parts.append(f"**我追求的**：{pursue or '（未填写）'}\n")

    # ── 二、交叉分析（复用 step2 引擎） ──
    tags_good = _scan_tags(good_at)
    tags_like = _scan_tags(like)
    tags_pursue = _scan_tags(pursue)
    cross_good_like = tags_good & tags_like
    cross_like_pursue = tags_like & tags_pursue
    cross_good_pursue = tags_good & tags_pursue
    cross_all = tags_good & tags_like & tags_pursue

    parts.append("## 二、交叉分析\n")
    if cross_all:
        parts.append(f"你的三圆交集（擅长 ∩ 喜欢 ∩ 追求）：**{'、'.join(cross_all)}**——这是你最核心的竞争力方向。\n")
    if cross_good_like - cross_all:
        parts.append(f"擅长 ∩ 喜欢：**{'、'.join(cross_good_like - cross_all)}**——你做得好且乐在其中的领域，最容易进入心流状态。")
    if cross_like_pursue - cross_all:
        parts.append(f"喜欢 ∩ 追求：**{'、'.join(cross_like_pursue - cross_all)}**——有热情且符合价值观，值得长期投入但需补齐能力。")
    if cross_good_pursue - cross_all:
        parts.append(f"擅长 ∩ 追求：**{'、'.join(cross_good_pursue - cross_all)}**——有能力且回报匹配，但注意是否真心喜欢，避免倦怠。")
    if not (cross_good_like or cross_like_pursue or cross_good_pursue):
        parts.append("三个维度暂未发现明显交集，建议拓宽关键词描述，或尝试新领域找到交叉点。")
    parts.append("")

    # ── 三、聚焦方向 ──
    parts.append("## 三、聚焦方向\n")
    if focus:
        parts.append(f"你确认的聚焦方向：**{focus}**\n")
        # 用方向规则引擎评估 focus 与标签的吻合度
        focus_lower = focus.lower()
        matched_dir = None
        for rule in _DIRECTION_RULES:
            if any(kw in focus_lower for kw in rule["name"].lower().replace("/", " ").split()):
                matched_dir = rule
                break
        if matched_dir:
            # 检查用户标签与该方向所需标签的重合
            dir_tags = set()
            for pair in matched_dir["require_any"]:
                dir_tags.update(pair)
            overlap = (tags_good | tags_like | tags_pursue) & dir_tags
            if overlap:
                parts.append(f"你的关键词中，**{'、'.join(overlap)}** 与该方向高度相关，方向选择有充分依据。")
            else:
                parts.append(f"你的关键词与该方向的核心要求（{'、'.join(list(dir_tags)[:4])}）重合度偏低，建议补充相关经历或重新审视。")
    else:
        parts.append("（未确认聚焦方向）\n")

    # ── 四、目标岗位评估 ──
    parts.append("\n## 四、目标岗位\n")
    if target_jobs:
        parts.append(f"你选择了 {len(target_jobs)} 个目标岗位：**{'、'.join(target_jobs)}**\n")
        # 对每个岗位做简要评估
        all_tags = tags_good | tags_like | tags_pursue
        for job_name in target_jobs[:5]:
            pos = POS_BY_ID.get(job_name) or next((p for p in POSITIONS if p["name"] == job_name), None)
            if not pos:
                continue
            job_skills = set(pos.get("skills") or [])
            tag_skill_overlap = all_tags & {t.lower() for t in job_skills}
            salary = pos.get("salary_range", [0, 0])
            salary_str = f"{salary[0]}-{salary[1]}元/月" if salary[0] > 0 else "待探索"
            cities = "、".join((pos.get("cities") or [])[:3]) or "不限"
            parts.append(f"- **{pos['name']}**（{cities} / {salary_str}）")
            if tag_skill_overlap:
                parts.append(f"  你的 {'、'.join(tag_skill_overlap)} 标签与该岗位直接相关")
            elif job_skills:
                parts.append(f"  该岗位核心要求：{'、'.join(list(job_skills)[:4])}，建议针对性补强")
    else:
        parts.append("（未选择目标岗位）\n")

    # ── 五、综合建议 ──
    parts.append("\n## 综合建议\n")

    # 方向推荐（复用 _DIRECTION_RULES 引擎）
    scored_dirs: list[tuple[str, float, str, list[str]]] = []
    for rule in _DIRECTION_RULES:
        score = 0.0
        for pair in rule["require_any"]:
            t0, t1 = pair
            dims0 = sum([t0 in tags_good, t0 in tags_like, t0 in tags_pursue])
            dims1 = sum([t1 in tags_good, t1 in tags_like, t1 in tags_pursue])
            if dims0 > 0 and dims1 > 0:
                score += 2.0 if (dims0 + dims1 >= 3) else 1.0
            elif dims0 > 0 or dims1 > 0:
                score += 0.3
        if score >= 1.0:
            stars = "★" * min(5, int(score))
            scored_dirs.append((rule["name"], score, rule["reason"], rule["jobs"]))
    scored_dirs.sort(key=lambda x: -x[1])

    if scored_dirs:
        parts.append("根据你的三维度交叉分析，最匹配的职业方向：\n")
        for name, sc, reason, jobs in scored_dirs[:3]:
            stars = "★" * min(5, max(1, int(sc)))
            parts.append(f"**{name}** {stars}")
            parts.append(f"  {reason}")
            parts.append(f"  推荐岗位：{'、'.join(jobs[:4])}\n")

    # 行动建议
    parts.append("### 下一步行动\n")
    if focus and target_jobs:
        parts.append(f"1. 用「就业倒推法」调研：{target_jobs[0]} 的中位数毕业生实际去了哪里、做什么、薪资多少")
        parts.append(f"2. 从 {len(target_jobs)} 个目标岗位中选 1-2 个，精读真实 JD，标注核心技能要求")
        parts.append("3. 用「500 强测试」验证：看头部企业实际在招这些岗位吗？要求什么？")
        parts.append("4. 寻找相关实习或项目机会，用实践验证方向是否适合")
        parts.append("5. 思考你的「不可替代性」：你的哪些能力组合是别人难以复制的？")
    elif focus:
        parts.append(f"1. 在「{focus}」方向下浏览具体岗位，用「中位数原则」关注大多数人的真实去向")
        parts.append("2. 与该方向的从业者交流（学长学姐、行业社群），获取一手信息")
        parts.append("3. 尝试小项目或实习，用行动验证兴趣")
        parts.append("4. 用「10 年后压迫测试」评估：这个方向长期还有竞争力吗？")
    else:
        parts.append("1. 回顾「擅长」和「喜欢」的交叉区域，找到你既能做好又乐在其中的方向")
        parts.append("2. 利用学校资源（职业中心、校友网络）做 2-3 次职业访谈")
        parts.append("3. 记住：选择 > 努力，方向错误的努力是浪费")
        parts.append("4. 选定方向后回来完成后续步骤")

    # 张雪峰视角
    zxf_tags = list(tags_good | tags_like | tags_pursue)[:5]
    zxf = _zxf_by_tags(zxf_tags, 2)
    if not zxf and target_jobs:
        zxf = _zxf_by_job(target_jobs[0], 2)
    if zxf:
        parts.append("\n### 张雪峰说\n")
        for q in zxf:
            parts.append(f"> 「{q['text']}」")
            if q.get("context"):
                parts.append(f"> *—— {q['context']}*")

    return "\n".join(parts)


@app.post("/api/assessment/deep/analyze-step2")
def analyze_step2(body: Step2In, request: Request):
    """Step2: 从 Step1 的三个文本 + 测评结果中分析交叉方向"""
    user = _current_user(request)
    assessments = store.get(user).get("assessments", {})

    good_at = (body.good_at or "").strip()
    like = (body.like or "").strip()
    pursue = (body.pursue or "").strip()

    if not good_at and not like and not pursue:
        raise HTTPException(400, "请至少填写一个维度")

    # 尝试 AI 分析
    user_text = (
        f"【我擅长的】\n{good_at}\n\n"
        f"【我喜欢的】\n{like}\n\n"
        f"【我追求的】\n{pursue}"
    )

    # 附加测评信息
    extra = ""
    holland = assessments.get("holland")
    mbti = assessments.get("mbti")
    if holland:
        order = ["R", "I", "A", "S", "E", "C"]
        labels = {"R": "实际型", "I": "研究型", "A": "艺术型", "S": "社会型", "E": "企业型", "C": "常规型"}
        ranked = sorted(range(len(order)), key=lambda i: (holland[i] if i < len(holland) else 0), reverse=True)
        top = [labels.get(order[i], "") for i in ranked[:3]]
        extra += f"\n\n用户的霍兰德测评 Top3：{'、'.join(top)}"
    if mbti:
        pairs = [("E", "I"), ("S", "N"), ("T", "F"), ("J", "P")]
        code = ""
        for i, (a, b) in enumerate(pairs):
            v = mbti[i] if i < len(mbti) else 0.5
            code += b if v >= 0.5 else a
        extra += f"\n用户的 MBTI：{code}"

    offline = _offline_step2_analysis(good_at, like, pursue, assessments)

    ai_result = _ai_or(
        "你是一位专业的职业规划顾问。用户正在做「思索工作的步骤」的头脑风暴，"
        "请分析用户填写的三个维度（擅长的/喜欢的/追求的），找出它们的交叉点和聚焦方向。\n"
        "要求：\n"
        "1) 用 Markdown 格式，200-400 字\n"
        "2) 先逐个总结三个维度的关键信息\n"
        "3) 找出交叉区域（擅长∩喜欢、喜欢∩追求、擅长∩追求）\n"
        "4) 给出 1-3 个建议的聚焦方向（具体、可执行）\n"
        "5) 如有测评数据，结合测评结论增强分析\n"
        "6) 语言简洁专业，直接给结论" + extra,
        user_text,
        offline["analysis"],
    )

    return {
        "analysis": ai_result or offline["analysis"],
        "focus_suggestion": offline["focus_suggestion"],
        "directions": offline["directions"],
        "recommended_jobs": offline.get("recommended_jobs", []),
    }


@app.post("/api/assessment/deep")
def deep_assess(body: DeepAssessIn, request: Request):
    answers = body.answers or {}

    good_at = (answers.get("good_at") or "").strip()
    like = (answers.get("like") or "").strip()
    pursue = (answers.get("pursue") or "").strip()
    if not good_at and not like and not pursue:
        raise HTTPException(400, "请至少完成第一步「头脑风暴」")

    user = _current_user(request)
    assessments = store.get(user).get("assessments", {})

    user_text = (
        f"【我擅长的】\n{good_at}\n\n"
        f"【我喜欢的】\n{like}\n\n"
        f"【我追求的】\n{pursue}\n\n"
        f"【聚焦方向】\n{answers.get('focus', '（未填写）')}\n\n"
        f"【目标岗位】\n{', '.join(answers.get('target_jobs', []))}\n\n"
        f"【过滤条件】\n{json.dumps(answers.get('conditions', {}), ensure_ascii=False)}"
    )

    extra = ""
    holland = assessments.get("holland")
    mbti = assessments.get("mbti")
    if holland:
        order = ["R", "I", "A", "S", "E", "C"]
        labels = {"R": "实际型", "I": "研究型", "A": "艺术型", "S": "社会型", "E": "企业型", "C": "常规型"}
        ranked = sorted(range(len(order)), key=lambda i: (holland[i] if i < len(holland) else 0), reverse=True)
        top = [labels.get(order[i], "") for i in ranked[:3]]
        extra += f"\n\n用户的霍兰德测评 Top3：{'、'.join(top)}"
    if mbti:
        pairs = [("E", "I"), ("S", "N"), ("T", "F"), ("J", "P")]
        code = ""
        for i, (a, b) in enumerate(pairs):
            v = mbti[i] if i < len(mbti) else 0.5
            code += b if v >= 0.5 else a
        extra += f"\n用户的 MBTI：{code}"

    analysis = _ai_or(
        "你是一位融合心理学与生涯规划的资深顾问。用户完成了「思索工作的步骤」四步探索，"
        "请基于用户的输入生成一份完整的职业方向分析报告（Markdown 格式，400-600 字）。\n"
        "要求：\n"
        "1) 从头脑风暴（擅长/喜欢/追求）出发，分析交叉点\n"
        "2) 评估用户确认的聚焦方向是否合理\n"
        "3) 对用户选择的目标岗位逐一点评（适合度、发展前景）\n"
        "4) 结合过滤条件给出最终的 Top 推荐\n"
        "5) 如有测评数据，结合 Holland/MBTI 结论增强分析\n"
        "6) 给出 2-3 条具体可执行的下一步行动建议\n"
        "7) 语言要温暖专业，避免心灵鸡汤" + extra,
        user_text,
        _offline_deep_analysis(answers, assessments),
    )

    deep_record = {"answers": answers, "analysis": analysis}
    store.patch("assessments.deep", deep_record, user=user)
    return deep_record


@app.get("/api/assessment/deep")
def get_deep_assess(request: Request):
    return store.get(_current_user(request)).get("assessments", {}).get("deep") or {}


# ---------- positions ----------
@app.get("/api/positions")
def list_positions(q: str = ""):
    items = POSITIONS
    if q:
        items = [p for p in items if q.lower() in p["name"].lower()]
    # lightweight list
    return [{
        "id": p["id"], "name": p["name"], "industries": p["industries"],
        "cities": p["cities"], "salary_range": p["salary_range"],
        "skills": p["skills"][:6], "tier": p["tier"],
    } for p in items]


@app.get("/api/positions/{pid}")
def get_position(pid: str):
    p = POS_BY_ID.get(pid)
    if not p:
        raise HTTPException(404, "not found")
    return p


TYPICAL_DAY_FAMILIES = [
    {
        "key": "dev",
        "match": ["java", "c/c++", "c++", "前端", "后端", "全栈", "app", "android", "ios", "小程序"],
        "excl": ["测试", "推广", "销售", "硬件"],
        "events": [
            ("09:30", "到岗，拉最新代码，查看邮件与 Jira"),
            ("10:00", "晨会：同步需求进度与阻塞点"),
            ("10:30", "拉分支，编写核心功能代码"),
            ("12:30", "午餐 & 小憩"),
            ("14:00", "代码评审（Code Review）与单元测试"),
            ("15:30", "与前端/后端/测试同学联调接口"),
            ("17:00", "修复今日 Bug，提交 PR / MR"),
            ("18:30", "日报：记录进度，规划明日任务"),
        ],
    },
    {
        "key": "test",
        "match": ["测试工程师", "软件测试", "硬件测试", "质量管理", "质检"],
        "events": [
            ("09:00", "到岗，同步昨日缺陷处理状态"),
            ("09:30", "晨会：与开发确认当日新需求与优先级"),
            ("10:00", "设计 / 执行测试用例"),
            ("11:30", "提交 Bug 到缺陷管理系统"),
            ("12:30", "午餐"),
            ("14:00", "回归测试与自动化脚本维护"),
            ("16:00", "与开发联调，验证修复结果"),
            ("18:00", "输出今日测试报告与质量数据"),
        ],
    },
    {
        "key": "sales",
        "match": ["销售", "广告销售", "电话销售", "网络销售", "大客户", "销售工程师", "销售运营", "销售助理", "广告", "bd", "商务"],
        "events": [
            ("08:30", "到岗，查看昨日订单与待跟进客户"),
            ("09:00", "晨会：业绩通报 + 今日目标拆解"),
            ("09:30", "电话 / 企微跟进意向客户"),
            ("11:00", "客户拜访或线上 Demo 演示"),
            ("12:30", "午餐（常为客户午餐）"),
            ("13:30", "继续外呼 / 拜访跟进"),
            ("16:00", "商务谈判、合同处理"),
            ("18:00", "录入 CRM，复盘转化数据"),
            ("19:00", "1v1 辅导或产品培训"),
        ],
    },
    {
        "key": "service",
        "match": ["客服", "售后", "网络客服", "电话客服", "呼叫", "内容审核"],
        "events": [
            ("08:30", "打开工单系统，早班打卡"),
            ("09:00", "接听呼入 / 在线消息响应"),
            ("11:00", "处理升级工单，协调技术支持"),
            ("12:00", "轮岗午餐（线路不断）"),
            ("13:30", "继续接线 / 回访老客户"),
            ("15:00", "案例归档与知识库更新"),
            ("17:00", "交接班，输出服务日志"),
        ],
    },
    {
        "key": "product",
        "match": ["产品专员", "产品经理", "产品助理"],
        "events": [
            ("09:30", "查看数据看板与用户反馈"),
            ("10:00", "需求评审会"),
            ("11:00", "画原型 / 撰写 PRD"),
            ("12:30", "午餐"),
            ("14:00", "与设计 / 开发 / 测试对齐细节"),
            ("15:30", "竞品调研 & 用户访谈"),
            ("17:00", "输出需求文档，提交下轮评审"),
            ("18:30", "日总结"),
        ],
    },
    {
        "key": "operations",
        "match": ["运营", "社区运营", "游戏运营", "运营助理", "销售运营", "推广", "app推广", "游戏推广"],
        "events": [
            ("09:30", "查看昨日 DAU / GMV / UGC 数据"),
            ("10:00", "内容排期会 / 活动脑暴"),
            ("11:00", "撰写推文 / 制作活动素材"),
            ("12:30", "午餐"),
            ("14:00", "发布内容，监测用户反馈"),
            ("15:30", "与客服 / 市场联动处理投诉"),
            ("17:00", "数据复盘，调整推广策略"),
            ("19:00", "夜间活动监播（视情况）"),
        ],
    },
    {
        "key": "hr",
        "match": ["招聘", "猎头", "培训师"],
        "events": [
            ("09:00", "查看候选人邮件、猎聘后台"),
            ("09:30", "电话沟通意向候选人"),
            ("11:00", "简历筛选 + 约面"),
            ("12:30", "午餐"),
            ("14:00", "组织现场 / 视频面试"),
            ("16:00", "Offer 谈判与背调"),
            ("18:00", "录入 ATS 系统，更新招聘漏斗"),
        ],
    },
    {
        "key": "legal",
        "match": ["律师", "法务", "知识产权", "专利"],
        "events": [
            ("09:00", "查阅最新案件 / 卷宗"),
            ("10:00", "与客户 / 当事人沟通"),
            ("11:30", "撰写合同、法律意见书"),
            ("12:30", "午餐"),
            ("14:00", "案件研究、法规检索"),
            ("15:30", "出庭或谈判（视安排）"),
            ("18:00", "总结今日工作，准备明日材料"),
        ],
    },
    {
        "key": "consulting",
        "match": ["咨询顾问", "顾问"],
        "events": [
            ("09:00", "项目晨会：目标与任务分工"),
            ("09:30", "案头研究：行业报告 / 数据分析"),
            ("11:30", "与客户访谈"),
            ("12:30", "午餐"),
            ("14:00", "搭建 PPT 框架，产出中间交付物"),
            ("16:30", "内部评审：与经理对齐思路"),
            ("19:00", "继续打磨交付物，夜间发送给客户"),
        ],
    },
    {
        "key": "admin",
        "match": ["档案", "资料", "统计员", "总助", "ceo", "董事长", "助理", "行政", "邮件员"],
        "excl": ["律师助理", "产品助理", "销售助理", "运营助理", "项目助理"],
        "events": [
            ("08:30", "到岗，处理邮件与会议安排"),
            ("09:00", "领导行程协调，准备会议材料"),
            ("10:00", "文档归档、数据录入"),
            ("12:00", "午餐"),
            ("13:30", "接待来访、会议纪要"),
            ("15:00", "报销、物资采购等事务处理"),
            ("17:00", "当日事项清单，明日准备"),
        ],
    },
    {
        "key": "project",
        "match": ["项目经理", "项目主管", "项目专员", "项目助理", "项目招投标", "bd经理", "储备", "管培生"],
        "events": [
            ("09:00", "晨会：各组进度同步"),
            ("10:00", "风险梳理、项目计划更新"),
            ("11:00", "与客户 / 供应商对接"),
            ("12:30", "午餐"),
            ("14:00", "跨部门协调会"),
            ("16:00", "撰写项目文档 / 招投标材料"),
            ("18:00", "日总结、问题升级"),
        ],
    },
    {
        "key": "translate",
        "match": ["翻译", "英语翻译", "日语翻译"],
        "events": [
            ("09:00", "查看今日翻译任务与参考术语"),
            ("09:30", "文档初译"),
            ("12:00", "午餐"),
            ("13:30", "继续翻译 / 会议口译（视需要）"),
            ("15:30", "术语整理、自校润色"),
            ("17:00", "交付审校，回复客户反馈"),
        ],
    },
    {
        "key": "research",
        "match": ["科研", "研究员"],
        "events": [
            ("09:00", "实验设备 / 数据准备"),
            ("10:00", "组会：汇报本周进度"),
            ("11:00", "文献阅读与笔记"),
            ("12:30", "午餐"),
            ("14:00", "实验 / 数据分析"),
            ("16:30", "整理实验记录"),
            ("18:00", "撰写报告 / 论文"),
        ],
    },
    {
        "key": "field",
        "match": ["风电", "实施工程师", "技术支持", "硬件"],
        "events": [
            ("08:00", "早会：安全交底与任务分工"),
            ("08:30", "前往现场 / 巡检设备"),
            ("10:00", "设备调试、故障排查"),
            ("12:00", "午餐（现场用餐）"),
            ("13:30", "继续作业 / 客户沟通"),
            ("16:00", "回公司整理数据与故障记录"),
            ("17:30", "输出维保报告，安全收工"),
        ],
    },
    {
        "key": "trainer",
        "match": ["培训师"],
        "events": [
            ("09:00", "备课：更新课件与案例"),
            ("10:00", "学员签到，课前互动"),
            ("10:30", "授课 / 实操指导"),
            ("12:00", "午餐"),
            ("13:30", "继续授课"),
            ("15:30", "个别辅导 + Q&A"),
            ("17:00", "作业批改、结训评估"),
        ],
    },
]

DEFAULT_FAMILY = {
    "key": "generic",
    "events": [
        ("09:00", "到岗，查看邮件与今日任务"),
        ("09:30", "部门晨会：同步进度与目标"),
        ("10:00", "核心业务处理"),
        ("12:00", "午餐"),
        ("14:00", "跨团队沟通与协作"),
        ("16:00", "完成当日产出"),
        ("18:00", "日报总结"),
    ],
}


def _classify_position(p: dict) -> dict:
    name = (p.get("name") or "").lower()
    for family in TYPICAL_DAY_FAMILIES:
        for kw in family["match"]:
            if kw.lower() in name:
                if any(x.lower() in name for x in family.get("excl", [])):
                    continue
                return family
    return DEFAULT_FAMILY


_NOISE_WORDS = (
    "知乎", "简书", "掘金", "CSDN", "脉脉", "腾讯网", "网易", "头条", "搜狐",
    "报名", "训练营", "限量", "内推", "优惠", "点击", "扫码", "关注公众号",
    "课程介绍", "官方证书", "立即下载", "官方网站", "收藏备用", "附模板",
    "推文内容", "推荐您搜索", "正式开启", "联合打造", "课程题目", "讲师介绍",
    "课程简介", "视频来咯", "人物介绍", "本期分享", "主题:", "第二十期",
    "http", "www.", ".com", "APP.", "√", "PDF",
)

# 只保留「像人话」的句子：必须含动词/助词/个人代词
_HUMAN_MARKERS = re.compile(
    r"(我|你|他|她|们|自己|大家|同事|老板|leader|团队|每天|日常|经常"
    r"|需要|应该|必须|可以|会|要|得|了|过|着|在|把|被|让|给|跟|和"
    r"|觉得|感觉|发现|认为|建议|希望|喜欢|讨厌|害怕|担心"
    r"|工作|上班|下班|加班|开会|沟通|汇报|写|做|学|看|改|调|跑|聊"
    r"|成长|提升|锻炼|收获|挑战|压力|成就|价值|意义)",
)


def _clean_snippets(raw: list[str], limit: int = 5) -> list[str]:
    """深度清洗搜索摘要：拆句 → 过滤噪音 → 去重 → 取最佳"""
    candidates: list[str] = []
    for text in raw:
        text = re.sub(r"<[^>]+>", "", text)
        # ── 去掉搜狗微信摘要中拼接的标题 ──
        # 格式1: "标题 正文" — 标题通常不含句号，后面是空格+正文
        text = re.sub(r"^【[^】]+】\s*", "", text)
        text = re.sub(r"^[^。！？\n]{4,40}[丨|｜][^。！？\n]{2,30}\s+", "", text)
        # 格式2: "标题...数字. 正文" 或 "标题 日期 正文"
        text = re.sub(r"^[^。！？\n]{4,50}\s+\d{4}年\d{1,2}月\d{1,2}日\s*·?\s*", "", text)
        # 格式3: 标题（无句号的短语）+ 空格 + 正文开始
        # "活用 AI 工作套路,产品经理每天多 50% 摸鱼时光 以及N个..."
        # → 找到第一个看起来像正文开始的位置（含人称/动词）
        m = re.match(r"^([^。！？\n]{8,50})\s{2,}", text)
        if m:
            text = text[m.end():]
        # 去末尾省略号
        text = re.sub(r"\.\.\.$", "", text)
        text = re.sub(r"…$", "", text)
        # 按句号/叹号/问号拆句
        for chunk in re.split(r"[。！？\n]", text):
            sent = chunk.strip()
            sent = re.sub(r"[.…]{2,}$", "", sent).strip()
            # 去掉搜索截断残留（以...开头或结尾的碎片）
            sent = re.sub(r"^\.\.\.\s*", "", sent).strip()
            sent = re.sub(r"\s*\.\.\.$", "", sent).strip()
            if len(sent) < 15 or len(sent) > 100:
                continue
            if any(w in sent for w in _NOISE_WORDS):
                continue
            if not _HUMAN_MARKERS.search(sent):
                continue
            # 排除文章标题拼接残留（含编号格式、或冒号结尾的标题）
            if re.match(r"^\d+[、.．]\s", sent):
                continue
            if re.match(r"^(第[一二三四五六七八九十]+[章节篇]|[一二三四五六七八九十]+、)", sent):
                continue
            candidates.append(sent)
    # 去重（完全相同 or 包含关系）
    unique: list[str] = []
    for c in candidates:
        if any(c in u or u in c for u in unique):
            continue
        unique.append(c)
    return unique[:limit]


def _typical_day_offline(p: dict) -> dict:
    family = _classify_position(p)
    events = [[t, e] for t, e in family["events"]]  # mutable copies
    name = p.get("name") or ""
    skills = p.get("skills") or []

    # ── 原有逻辑：注入岗位特色到第一个核心工作行 ──
    break_words = ("午餐", "晨会", "到岗", "日报", "总结", "交接班", "签到", "早会", "夜间", "晚班")
    core_idx = next(
        (i for i, (_, e) in enumerate(events) if not any(w in e for w in break_words)),
        0,
    )
    if skills:
        focus = "、".join(skills[:3])
        events[core_idx][1] += f"（主要使用：{focus}）"
    else:
        events[core_idx][1] += f"（聚焦 {name} 的核心业务）"

    # Sales / ops often end later; add a closing hint if position name hints 晚班
    if any(k in name for k in ["夜", "晚", "客服"]):
        events.append(["20:00", "交接晚班，记录值班日志"])

    keywords = skills[:8] if skills else [
        "协作", "交付", "复盘", "学习", "沟通", "目标", "执行", "反馈",
    ]

    result: dict = {
        "items": [{"time": t, "event": e} for t, e in events],
        "keywords": keywords,
        "family": family["key"],
    }

    # ── 附加从业者真实声音（来自搜索摘要） ──
    # 精确匹配 → 模糊匹配
    zhihu = ZHIHU_EXP.get(name, {})
    if not zhihu:
        for zk, zv in ZHIHU_EXP.items():
            if zk in name or name in zk:
                zhihu = zv
                break

    if zhihu.get("daily_snippets"):
        result["daily_snippets"] = _clean_snippets(zhihu["daily_snippets"], 5)
    if zhihu.get("ability_insights"):
        result["ability_insights"] = _clean_snippets(zhihu["ability_insights"], 5)
    if zhihu.get("feelings"):
        result["feelings"] = _clean_snippets(zhihu["feelings"], 5)

    # 张雪峰视角（按岗位匹配相关语录）
    zxf = _zxf_by_job(name, 2)
    if zxf:
        result["zxf_quotes"] = [{"text": q["text"], "context": q.get("context", "")} for q in zxf]

    return result


@app.get("/api/positions/{pid}/typical-day")
def typical_day(pid: str):
    p = POS_BY_ID.get(pid)
    if not p:
        raise HTTPException(404, "not found")

    # Only call the LLM if a real API key is configured; otherwise the
    # offline stub returns generic text that is worse than our deterministic
    # family-based template.
    if llm.API_KEY:
        prompt = (
            f"请为岗位 '{p['name']}' 生成一张典型的一天日程表。要求：\n"
            f"1) 时间要贴合该岗位真实的上下班节奏（研发 9:30-19:00，销售 8:30-19:00，"
            f"客服轮班，现场工程早 8:00 等）\n"
            f"2) 事件要融入该岗位常用的核心技能/工具：{', '.join(p.get('skills', [])[:6])}\n"
            f"3) 每条日程要具体，避免出现 '核心工作' '完成产出' 这类空话\n"
            f"4) 共 6-9 条\n"
            f"输出 JSON：{{items:[{{time,event}}], keywords:[]}}\n\n"
            f"职位描述：{p.get('description', '')[:500]}"
        )
        data = llm.chat_json(
            "你是熟悉各行各业真实工作节奏的资深 HR，输出的日程表要具体可信，只输出 JSON。",
            prompt,
        )
        if isinstance(data, dict) and "items" in data and data["items"]:
            return data

    return _typical_day_offline(p)


@app.get("/api/graph")
def full_graph():
    return GRAPH


@app.get("/api/positions/{pid}/graph")
def position_graph(pid: str):
    """Return the promotion + switch paths centered on this position."""
    promote_nodes: set[str] = set()
    switch_nodes: set[str] = set()
    promote_edges = []
    switch_edges = []
    # build forward closure
    outgoing_promote = [e for e in GRAPH["edges"] if e["type"] == "promote"]

    def bfs(start: str):
        seen = {start}
        frontier = [start]
        while frontier:
            nxt = []
            for node in frontier:
                for e in outgoing_promote:
                    if e["source"] == node and e["target"] not in seen:
                        seen.add(e["target"])
                        nxt.append(e["target"])
                        promote_edges.append(e)
                        promote_nodes.add(e["target"])
                        promote_nodes.add(e["source"])
            frontier = nxt

    bfs(pid)
    for e in GRAPH["edges"]:
        if e["type"] == "switch" and e["source"] == pid:
            switch_edges.append(e)
            switch_nodes.add(e["source"])
            switch_nodes.add(e["target"])

    nodes = {n["id"]: n for n in GRAPH["nodes"]}
    return {
        "promote": {
            "nodes": [nodes[n] for n in promote_nodes if n in nodes] or [nodes[pid]],
            "edges": promote_edges,
        },
        "switch": {
            "nodes": [nodes[n] for n in switch_nodes if n in nodes] or [nodes[pid]],
            "edges": switch_edges,
        },
    }


# ---------- matching ----------
def _student_for_match(state: dict) -> dict:
    """Bundle profile + assessments so matching.match() can see interest scores."""
    prof = dict(state.get("profile") or {})
    prof["assessments"] = state.get("assessments") or {}
    return prof


@app.get("/api/match/recommend")
def recommend(request: Request, limit: int = 10):
    s = store.get(_current_user(request))
    student = _student_for_match(s)
    scored = []
    for p in POSITIONS:
        m = matching.match(student, p)
        hit = m.get("skills_hit", [])
        reason = f"匹配技能：{'、'.join(hit[:3])}" if hit else "综合素质匹配"
        scored.append({
            "id": p["id"], "name": p["name"],
            "industries": p["industries"], "cities": p["cities"],
            "salary_range": p["salary_range"], "skills": p["skills"][:6],
            "match": m["score"],
            "reason": reason,
        })
    scored.sort(key=lambda x: x["match"], reverse=True)
    return scored[:limit]


def _match_advice(m: dict, job: dict) -> str:
    """基于匹配结果生成文字解读"""
    if m.get("empty_profile"):
        return "请先完善个人档案，系统才能为你提供准确的匹配分析。"
    parts: list[str] = []
    hit = m.get("skills_hit", [])
    miss = m.get("skills_miss", [])
    bd = m.get("breakdown", {})
    score = m.get("score", 0)
    name = job.get("name", "")

    # 总体评价
    if score >= 80:
        parts.append(f"你与「{name}」高度匹配，综合竞争力突出。")
    elif score >= 60:
        parts.append(f"你与「{name}」匹配度良好，有针对性提升后竞争力会更强。")
    elif score >= 40:
        parts.append(f"你与「{name}」有一定基础，但需要较多准备。")
    else:
        parts.append(f"你与「{name}」当前匹配度偏低，建议先了解岗位要求再做规划。")

    # 优势
    if hit:
        parts.append(f"你的 {'、'.join(hit[:4])} 技能与该岗位直接匹配，这是你的核心优势。")
    if bd.get("职业素养", 0) >= 70:
        parts.append("你的职业素养评分较高，说明软技能和兴趣方向与该岗位契合。")

    # 短板
    if miss:
        parts.append(f"建议补强：{'、'.join(miss[:4])}——可通过在线课程、项目实践或实习来积累。")
    if bd.get("基础要求", 0) < 50:
        det = m.get("details", {})
        if det.get("专业", {}).get("state") == "mismatch":
            parts.append("专业背景与该岗位不完全对口，建议通过辅修、证书或相关项目弥补。")

    return "".join(parts)


@app.get("/api/match/{pid}")
def match_one(pid: str, request: Request):
    p = POS_BY_ID.get(pid)
    if not p:
        raise HTTPException(404, "not found")
    s = store.get(_current_user(request))
    m = matching.match(_student_for_match(s), p)
    m["advice"] = _match_advice(m, p)
    return m


class BatchMatchIn(BaseModel):
    job_names: list[str]


@app.post("/api/match/batch")
def match_batch(body: BatchMatchIn, request: Request):
    """批量匹配：根据岗位名模糊查找 positions，返回市场数据 + 匹配分数"""
    s = store.get(_current_user(request))
    student = _student_for_match(s)
    results = []
    seen = set()

    for name in body.job_names:
        name = name.strip()
        if not name:
            continue
        # 精确匹配 → 模糊包含
        p = POS_BY_ID.get(name)
        if not p:
            for pos in POSITIONS:
                if name in pos["name"] or pos["name"] in name:
                    p = pos
                    break
        if not p or p["id"] in seen:
            # 找不到的岗位也返回，标记 found=false
            if name not in seen:
                results.append({"name": name, "found": False})
                seen.add(name)
            continue
        seen.add(p["id"])
        m = matching.match(student, p)
        results.append({
            "name": p["name"],
            "id": p["id"],
            "found": True,
            "salary_range": p["salary_range"],
            "cities": p["cities"],
            "skills": p["skills"][:8],
            "education": p.get("education", ""),
            "industries": p["industries"],
            "work_env": p.get("work_env", ""),
            "tier": p["tier"],
            "count": p.get("count", 0),
            "match_score": m["score"],
            "breakdown": m.get("breakdown", {}),
            "skills_hit": m.get("skills_hit", []),
            "skills_miss": m.get("skills_miss", []),
        })

    return results


# ---------- report (Venn + plan) ----------
class ReportIn(BaseModel):
    target_position_id: str | None = None  # None = 使用推荐 Top 1


def _format_experiences(prof: dict) -> str:
    """把学生的在校/实习/工作经历压成便于 LLM 消化的文本。"""
    lines: list[str] = []

    eduction = prof.get("education") or []
    if isinstance(eduction, list) and eduction:
        lines.append("【在校经历】")
        for e in eduction:
            if not isinstance(e, dict):
                continue
            parts = [
                " ".join(x for x in [e.get("period_start"), "~", e.get("period_end") or "至今"] if x),
                " ".join(x for x in [e.get("school"), e.get("degree"), e.get("major")] if x),
                f"学院：{e['college']}" if e.get("college") else "",
                f"方向：{e['direction']}" if e.get("direction") else "",
                f"导师：{e['advisor']}" if e.get("advisor") else "",
            ]
            line = " · ".join(p for p in parts if p)
            if line:
                lines.append("- " + line)

    internships = prof.get("internships") or []
    if internships:
        lines.append("【实习经历】")
        for e in internships:
            if not isinstance(e, dict):
                continue
            head = " · ".join(x for x in [e.get("company"), e.get("position")] if x)
            period = " ~ ".join(x for x in [e.get("period_start"), e.get("period_end") or "至今"] if x)
            desc = (e.get("desc") or "").strip()
            line = f"- {head}"
            if period:
                line += f"（{period}）"
            if desc:
                line += "：" + desc[:300]
            lines.append(line)

    work = prof.get("work_experiences") or []
    if work:
        lines.append("【工作经历】")
        for e in work:
            if not isinstance(e, dict):
                continue
            head = " · ".join(x for x in [e.get("company"), e.get("position")] if x)
            period = " ~ ".join(x for x in [e.get("period_start"), e.get("period_end") or "至今"] if x)
            desc = (e.get("desc") or "").strip()
            line = f"- {head}"
            if period:
                line += f"（{period}）"
            if desc:
                line += "：" + desc[:300]
            lines.append(line)

    projects = prof.get("projects") or []
    if projects:
        lines.append("【项目经历】")
        for e in projects:
            if not isinstance(e, dict):
                continue
            head = " · ".join(x for x in [e.get("name"), e.get("role")] if x)
            period = " ~ ".join(x for x in [e.get("period_start"), e.get("period_end") or "至今"] if x)
            desc = (e.get("desc") or "").strip()
            line = f"- {head}"
            if period:
                line += f"（{period}）"
            if desc:
                line += "：" + desc[:300]
            lines.append(line)

    skills = prof.get("skills") or []
    if skills:
        lines.append("【专业技能】" + "、".join(skills[:15]))

    return "\n".join(lines) if lines else "（学生尚未填写任何经历）"


def _format_intent(prof: dict) -> str:
    intent = prof.get("intent") or {}
    if not intent:
        return "（学生尚未填写求职意向）"
    # 多城市兼容：优先 cities 数组，回退到 city
    cities_list = intent.get("cities") or []
    if isinstance(cities_list, list) and cities_list:
        city_text = "、".join(cities_list)
    else:
        city_text = intent.get("city") or ""
    # 工作时间：上班 - 下班 + 每周天数
    ws = intent.get("work_start") or ""
    we = intent.get("work_end") or ""
    daily = f"{ws}-{we}" if ws and we else (ws or we)
    wd = intent.get("work_days") or ""
    legacy_wh = intent.get("work_hours") or ""
    legacy_wt = intent.get("work_time") or ""
    work_text = " · ".join(x for x in [daily, wd, legacy_wh, legacy_wt] if x)
    parts = [
        f"意向行业：{intent.get('industry')}" if intent.get("industry") else "",
        f"意向岗位：{intent.get('job')}" if intent.get("job") else "",
        f"意向城市：{city_text}" if city_text else "",
        f"期望月薪：{intent.get('salary_min')} 元" if intent.get("salary_min") else "",
        f"工作时间：{work_text}" if work_text else "",
        f"工作环境：{intent.get('env')}" if intent.get("env") else "",
    ]
    return "\n".join(p for p in parts if p) or "（学生尚未填写求职意向）"


def _format_assessments(assess: dict) -> str:
    """Flatten the four self-assessment results into LLM-friendly text."""
    if not assess:
        return "（学生尚未完成任何自我认识测评）"
    lines: list[str] = []

    holland = assess.get("holland") or []
    if holland and len(holland) >= 6:
        names = HOLLAND_NAMES  # 已在 interpret_assess 区域定义
        pairs = sorted(zip(names, holland), key=lambda x: -x[1])
        code = _holland_code(holland)
        top3 = "、".join(f"{n}({s}分)" for n, s in pairs[:3])
        lines.append(f"【霍兰德职业兴趣】代码 {code}；Top 3：{top3}")

    mi = assess.get("multi_intel") or []
    if mi and len(mi) >= 6:
        pairs = sorted(zip(MI_NAMES, mi), key=lambda x: -x[1])
        top3 = "、".join(f"{n}({s}分)" for n, s in pairs[:3])
        lines.append(f"【多元智能】Top 3 优势智能：{top3}")

    values = assess.get("values") or []
    if values and len(values) >= 6:
        pairs = sorted(zip(VALUE_NAMES, values), key=lambda x: -x[1])
        top3 = "、".join(f"{n}({s}分)" for n, s in pairs[:3])
        lines.append(f"【工作价值观】Top 3 核心驱动：{top3}")

    mbti = assess.get("mbti") or []
    mtype = _mbti_from_scores(mbti)
    if mtype:
        label = MBTI_DESCRIPTIONS.get(mtype, "")
        lines.append(f"【MBTI 人格类型】{mtype}" + (f"（{label}）" if label else ""))

    return "\n".join(lines) if lines else "（学生尚未完成任何自我认识测评）"


def _format_target_job(job: dict) -> str:
    """Flatten a position dict into a rich paragraph for the LLM."""
    lines = [f"岗位名称：{job.get('name', '')}"]
    if job.get("industries"):
        lines.append("行业：" + "、".join(job["industries"]))
    if job.get("cities"):
        lines.append("主要城市：" + "、".join(job["cities"][:5]))
    sr = job.get("salary_range") or [0, 0]
    if sr[1]:
        lines.append(f"薪资范围：{sr[0]}-{sr[1]} 元/月")
    if job.get("education"):
        lines.append(f"学历要求：{job['education']}")
    if job.get("skills"):
        lines.append("核心技能：" + "、".join(job["skills"][:10]))
    if job.get("certs"):
        lines.append("加分证书：" + "、".join(job["certs"]))
    soft = job.get("soft_skills") or {}
    if soft:
        items = [f"{k} {v}" for k, v in list(soft.items())[:5]]
        lines.append("软技能画像：" + "、".join(items))
    if job.get("description"):
        lines.append("\n岗位描述节选：\n" + str(job["description"])[:600])
    return "\n".join(lines)


def _ai_or(prompt_system: str, prompt_user: str, fallback: str) -> str:
    """调 LLM，返回空/offline-stub 时用 fallback。"""
    if not llm.API_KEY:
        return fallback
    out = llm.chat(prompt_system, prompt_user)
    if not out or "offline-stub" in out:
        return fallback
    return out


def _offline_experience_summary(prof: dict) -> str:
    edu = prof.get("education") or []
    intern = prof.get("internships") or []
    work = prof.get("work_experiences") or []
    projects = prof.get("projects") or []
    skills = prof.get("skills") or []

    parts: list[str] = []

    # 教育
    if edu:
        best = max(
            edu,
            key=lambda e: {"博士": 5, "硕士": 4, "本科": 3, "大专": 2, "高中": 1}.get(
                (e or {}).get("degree", ""), 0,
            ),
        )
        school = best.get("school") or "某高校"
        major = best.get("major") or ""
        degree = best.get("degree") or ""
        edu_line = f"你就读于 **{school}**"
        if degree:
            edu_line += f" {degree}"
        if major:
            edu_line += f"，主修 **{major}**"
        if best.get("direction"):
            edu_line += f"，专注于 {best['direction']} 方向"
        edu_line += "。"
        parts.append(edu_line)
    else:
        parts.append("你尚未填写教育背景，建议先在「我的档案」中补齐。")

    # 实习 / 工作 —— 提炼能力而非搬运描述
    all_jobs = list(intern) + list(work)
    if all_jobs:
        # 从所有经历描述中提取关键能力词
        all_desc = " ".join((e.get("desc") or "") for e in all_jobs if isinstance(e, dict))
        ability_map = {
            "数据分析": ["数据", "分析", "统计", "报表", "Excel", "SPSS", "SQL"],
            "内容创作": ["文案", "撰写", "编辑", "内容", "文章", "写作", "稿"],
            "用户研究": ["用户", "调研", "访谈", "问卷", "需求", "反馈"],
            "项目管理": ["项目", "协调", "推进", "跟进", "排期", "统筹"],
            "设计能力": ["设计", "海报", "UI", "视觉", "布局", "排版"],
            "运营推广": ["运营", "推广", "账号", "粉丝", "转化", "增长", "自媒体"],
            "技术开发": ["开发", "代码", "编程", "前端", "后端", "Python", "Java", "平台", "建站"],
            "沟通协作": ["沟通", "协作", "团队", "对接", "配合", "合作"],
            "心理咨询": ["咨询", "来访", "心理", "辅导", "个案"],
        }
        extracted: list[str] = []
        for ability, keywords in ability_map.items():
            if any(kw in all_desc for kw in keywords):
                extracted.append(ability)

        companies = [e.get("company") or "" for e in all_jobs if isinstance(e, dict) and e.get("company")]
        positions = [e.get("position") or "" for e in all_jobs if isinstance(e, dict) and e.get("position")]

        parts.append(
            f"**【实践经历】** 你累计有 **{len(all_jobs)}** 段实习/工作经历，"
            f"涉及 **{'、'.join(companies[:3])}** 等机构。"
        )
        if extracted:
            parts.append(
                f"从你的经历中提炼出的核心能力包括：**{'、'.join(extracted[:5])}**。"
                "这些能力是你区别于其他候选人的实战积累，建议在简历中用具体成果来佐证。"
            )
        if len(all_jobs) >= 2:
            parts.append(
                "多段经历说明你有一定的职场适应能力和探索意识，"
                "接下来的关键是找到经历之间的共同主线，形成你的职业叙事。"
            )
    else:
        parts.append("**【实践经历】** 你暂无实习或全职工作经历，建议尽快补齐 1-2 段与目标方向相关的实习。")

    # 项目 —— 提炼产出价值
    if projects:
        valid_projects = [p for p in projects if isinstance(p, dict) and (p.get("name") or p.get("desc"))]
        if valid_projects:
            proj_names = [str(p.get("name") or "项目") for p in valid_projects[:3]]
            # 从项目描述中提取关键词
            proj_desc_all = " ".join(str(p.get("desc") or "") for p in valid_projects)
            proj_tags = []
            for tag, kws in [("竞赛经验", ["竞赛", "大赛", "比赛"]), ("产品思维", ["产品", "用户", "需求"]),
                             ("数据驱动", ["数据", "分析", "模型"]), ("团队协作", ["团队", "协作", "合作"]),
                             ("技术实现", ["开发", "实现", "搭建", "代码"])]:
                if any(kw in proj_desc_all for kw in kws):
                    proj_tags.append(tag)
            parts.append(
                f"**【项目产出】** 你参与过 **{len(valid_projects)}** 个项目（{'、'.join(proj_names)}），"
                + (f"体现了你在 **{'、'.join(proj_tags[:3])}** 方面的能力。" if proj_tags else "这些都是展示你实际产出的有效素材。")
            )

    # 技能
    if skills:
        parts.append(
            f"**【技能储备】** 你目前在档案中填写了 {len(skills)} 项专业技能，核心技术栈为 **"
            + "、".join(skills[:6]) + "**。"
        )
        hard = [s for s in skills if any(kw in s.lower() for kw in ("python", "java", "c++", "sql", "excel", "spss", "r语言", "matlab", "git", "docker", "linux", "javascript", "typescript", "react", "vue", "node", "pytorch", "tensorflow", "数据分析", "机器学习"))]
        soft = [s for s in skills if any(kw in s for kw in ("沟通", "协作", "领导", "团队", "管理", "策划", "写作", "演讲", "谈判"))]
        if hard and soft:
            parts.append(
                f"其中硬技能（{' / '.join(hard[:4])}）构成你的专业门槛，"
                f"软技能（{' / '.join(soft[:3])}）则是你在团队中发挥价值的加速器。"
                "两者兼备是求职时的加分项，建议在简历中分开呈现、各举实例。"
            )
        elif hard:
            parts.append(
                f"你的技能偏向硬技能型（{' / '.join(hard[:4])}），技术基础扎实。"
                "建议补充 1-2 项软技能关键词（如团队协作、跨部门沟通），让简历更加立体。"
            )
    else:
        parts.append("**【技能储备】** 你尚未填写任何专业技能，建议补充 3-6 项核心技能标签以便匹配岗位。")

    # 综合评估
    total_exp = len(list(intern) + list(work))
    total_proj = len(projects)
    if total_exp >= 2 and total_proj >= 1:
        parts.append(
            f"**【综合评估】** 你已积累 {total_exp} 段实习/工作经历和 {total_proj} 个项目，"
            "在同届求职者中属于经验较丰富的梯队。接下来重点应放在'提炼成果'上 —— "
            "每段经历提炼 1-2 个可量化的成果（提升了 X%、节省了 Y 小时），"
            "这比单纯罗列职责更能打动面试官。"
        )
    elif total_exp >= 1:
        parts.append(
            f"**【综合评估】** 你目前有 {total_exp} 段经历，在同龄人中属于起步阶段。"
            "建议在毕业前再补充 1-2 段与目标方向强相关的实习，"
            "同时把现有经历中最亮眼的产出（数据、作品、用户反馈）整理成结构化案例。"
        )
    else:
        parts.append(
            "**【综合评估】** 你目前尚无实习或项目经历，这在求职中会是明显短板。"
            "建议优先找 1 段哪怕只有 2-3 个月的实习，或独立完成 1 个可展示的个人项目，"
            "用'做过'替代'学过'来建立你的职业起点。"
        )

    return "\n\n".join(parts)


HOLLAND_LETTER_NAME = {
    "R": "实际型", "I": "研究型", "A": "艺术型",
    "S": "社会型", "E": "企业型", "C": "常规型",
}


def _infer_holland_from_text(text: str) -> str:
    """粗略地把意向行业/岗位文本映射成霍兰德字母，用于对齐分析。"""
    if not text:
        return ""
    table = [
        (["研发", "开发", "算法", "数据", "测试", "科研", "分析", "AI", "程序员", "工程师", "咨询", "律师"], "I"),
        (["硬件", "实施", "运维", "风电", "机械", "电气", "网络"], "R"),
        (["销售", "BD", "商务", "运营", "推广", "产品", "项目经理", "管培", "创业", "市场"], "E"),
        (["客服", "招聘", "猎头", "培训", "教育", "社区", "用户研究", "HR", "人力"], "S"),
        (["设计", "UI", "UX", "内容", "编辑", "文案", "翻译", "品牌", "视觉", "艺术"], "A"),
        (["财务", "审计", "行政", "档案", "资料", "统计", "质检", "质量", "合规", "法务"], "C"),
    ]
    t = text.lower()
    for kws, letter in table:
        if any(kw.lower() in t for kw in kws):
            return letter
    return ""


# 价值观冲突检测（来自测评知识库 §3.3）
VALUE_CONFLICTS = [
    ("稳定安全", "自主独立", "你同时重视稳定与自由，需要在'体制保障'和'灵活自主'之间找到平衡点"),
    ("经济回报", "社会影响", "你既重视物质回报又看重社会价值，择业时可能需要在商业机会与公益价值间权衡"),
    ("学习成长", "生活平衡", "你希望快速成长又看重生活边界，高成长岗位往往需要更多时间投入，需评估可接受的边界"),
    ("稳定安全", "创造创新", "稳定与创新方向在多数环境下难以兼得，可以优先在创新性组织中寻找相对稳定的角色"),
]


def _offline_interest_summary(prof: dict, assess: dict | None = None) -> str:
    """基于档案 + 测评做跨维度推理，输出更深度的个人兴趣画像。"""
    intent = prof.get("intent") or {}
    assess = assess or {}

    # ---------- 收集原始数据 ----------
    holland = assess.get("holland") or []
    mi = assess.get("multi_intel") or []
    values = assess.get("values") or []
    mbti = assess.get("mbti") or []

    holland_pairs = (
        sorted(zip(HOLLAND_NAMES, holland), key=lambda x: -x[1])
        if holland and len(holland) >= 6 else []
    )
    holland_code = _holland_code(holland) if holland_pairs else ""
    holland_top_letter = holland_code[0] if holland_code else ""
    holland_top_name = holland_pairs[0][0] if holland_pairs else ""

    value_pairs = (
        sorted(zip(VALUE_NAMES, values), key=lambda x: -x[1])
        if values and len(values) >= 6 else []
    )
    value_top3 = [n for n, _ in value_pairs[:3]] if value_pairs else []
    value_bottom = value_pairs[-1][0] if value_pairs else ""
    value_highs = {n for n, s in value_pairs if s >= 70} if value_pairs else set()

    mi_pairs = (
        sorted(zip(MI_NAMES, mi), key=lambda x: -x[1])
        if mi and len(mi) >= 6 else []
    )
    mi_top3 = [n for n, _ in mi_pairs[:3]] if mi_pairs else []

    mtype = _mbti_from_scores(mbti)
    mbti_label = MBTI_DESCRIPTIONS.get(mtype, "")

    intent_industry = (intent.get("industry") or "").strip()
    intent_job = (intent.get("job") or "").strip()
    cities_list = intent.get("cities") or []
    if isinstance(cities_list, list) and cities_list:
        intent_city = "、".join(cities_list)
    else:
        intent_city = (intent.get("city") or "").strip()
    intent_salary = intent.get("salary_min") or 0
    _ws = intent.get("work_start") or ""
    _we = intent.get("work_end") or ""
    _daily = f"{_ws}-{_we}" if _ws and _we else (_ws or _we)
    intent_work = " · ".join(
        x for x in [
            _daily,
            intent.get("work_days") or "",
            intent.get("work_hours") or "",
            intent.get("work_time") or "",
        ] if x
    )
    intent_env = (intent.get("env") or "").strip()

    has_intent = any([intent_industry, intent_job, intent_city, intent_salary, intent_work, intent_env])
    has_assess = bool(holland_pairs or value_pairs or mi_pairs or mtype)

    paragraphs: list[str] = []

    # ---------- 段 1：融合画像（一句话定性） ----------
    lead_bits: list[str] = []
    if mtype and mbti_label:
        lead_bits.append(f"一个 **{mtype}·{mbti_label.split('·')[0].strip()}** 风格的个体")
    if holland_top_name:
        lead_bits.append(f"主导兴趣偏向 **{holland_top_name}**")
    if value_top3:
        lead_bits.append(f"最看重 **{' · '.join(value_top3)}**")
    if mi_top3:
        lead_bits.append(f"认知上擅长 **{' · '.join(mi_top3[:2])}**")

    if lead_bits:
        paragraphs.append(
            "**【融合画像】** 综合你的测评与意向数据，你是"
            + "，".join(lead_bits)
            + "。这个组合决定了你在职业选择中的底层节奏。"
        )
    elif has_intent:
        paragraphs.append(
            "**【融合画像】** 由于测评数据还不完整，以下分析主要基于你填写的求职意向。"
            "建议你先完成霍兰德和 MBTI 至少一项，系统会给出更精准的画像。"
        )
    else:
        paragraphs.append(
            "**【融合画像】** 你尚未完成任何测评，也尚未填写求职意向。"
            "请先去「我的档案」补齐意向，再到「自我认识」完成测评，AI 才能为你生成有意义的个人兴趣画像。"
        )

    # ---------- 段 2：求职意向 × 兴趣对齐 ----------
    if has_intent:
        intent_bits = []
        if intent_job:
            intent_bits.append(f"意向岗位 **{intent_job}**")
        if intent_city:
            intent_bits.append(f"期望城市 **{intent_city}**")
        if intent_salary:
            intent_bits.append(f"期望月薪 **{intent_salary} 元**以上")
        if intent_work:
            intent_bits.append(f"可接受工时 **{intent_work}**")

        intent_text = "你的求职意向包括：" + "、".join(intent_bits) + "。"

        implied_letter = _infer_holland_from_text(f"{intent_job} {intent_industry}")
        if implied_letter and holland_top_letter:
            implied_name = HOLLAND_LETTER_NAME.get(implied_letter, implied_letter)
            if implied_letter == holland_top_letter:
                intent_text += f" 该方向属于 **{implied_name}**，与你的主导兴趣一致。"
            elif implied_letter not in holland_code[:3]:
                intent_text += (
                    f" 值得注意的是，该方向属于 **{implied_name}**，"
                    f"而你的主导兴趣是 **{holland_top_letter} {holland_top_name}**，两者存在一定张力。"
                )

        paragraphs.append("**【意向对齐】** " + intent_text)

    # ---------- 段 3：认知优势 × 岗位能力 ----------
    if mi_top3 and intent_job:
        mi_text = f"你的优势智能集中在 **{' · '.join(mi_top3)}**，"
        # 典型职业能力组合映射
        combo_tips = [
            (("逻辑-数学", "自然观察"), "在算法、数据科学、结构化分析上会非常如鱼得水"),
            (("逻辑-数学", "空间"), "在前端、游戏、系统架构这些'既要抽象又要具象'的领域有天然优势"),
            (("空间", "语言"), "在 UI/UX 设计、内容产品、视觉叙事方向会脱颖而出"),
            (("人际", "语言"), "非常适合产品、用户研究、品牌/运营这类需要'理解人并把想法讲清楚'的岗位"),
            (("人际", "内省"), "适合 HR、生涯规划、心理咨询这类需要同时懂别人和懂自己的方向"),
            (("逻辑-数学", "人际"), "是 AI 产品经理、商业分析师这种稀缺复合型岗位的核心画像"),
            (("内省", "语言"), "在创业、内容创作、独立研究等需要自驱力 + 表达力的路径上会走得很远"),
            (("逻辑-数学", "身体-运动"), "适合嵌入式、运维、硬件这类'能动脑又能动手'的岗位"),
        ]
        matched_combo = ""
        mi_set = set(mi_top3)
        for combo, tip in combo_tips:
            if set(combo).issubset(mi_set):
                matched_combo = tip
                break
        if matched_combo:
            mi_text += matched_combo + "。"
        else:
            mi_text += f"这是你在 **{intent_job}** 方向可以持续输出的核心燃料。"
        paragraphs.append("**【认知优势】** " + mi_text)

    # ---------- 段 4：MBTI 工作风格 ----------
    if mtype and mbti_label:
        paragraphs.append(
            f"**【工作风格】** 你的 MBTI 是 **{mtype}**，对应的典型工作风格是 **{mbti_label}**。"
            "这是你最自然、最不消耗能量的做事方式，择业时要尽量保护这种风格不被组织文化消耗。"
        )

    # ---------- 段 5：冲突 / 盲区提醒 ----------
    alerts: list[str] = []
    for a, b, msg in VALUE_CONFLICTS:
        if a in value_highs and b in value_highs:
            alerts.append(msg)
    if value_bottom == "经济回报" and intent_salary and intent_salary >= 20000:
        alerts.append(
            f"经济回报是你价值观里最低的一项，但期望月薪设在 {intent_salary} 元，"
            "可能并非你真正的动机来源，需要明确为什么坚持这个数字。"
        )
    if holland_top_letter == "I" and intent_job and any(
        kw in intent_job for kw in ("销售", "BD", "商务")
    ):
        alerts.append(
            "你的主导兴趣是研究型，但意向岗位偏向销售/商务方向，这是一个明显的张力点，"
            "建议用一段短期实习验证你是否真的能长期从事高强度人际沟通的工作。"
        )
    if holland_top_letter == "E" and intent_job and any(
        kw in intent_job for kw in ("科研", "算法", "研究员", "后端")
    ):
        alerts.append(
            "你的主导兴趣是企业型，但意向岗位偏向深度研发/研究，"
            "这类岗位会限制你的影响力半径，想清楚是否接受长期的'幕后角色'。"
        )
    if mtype and mtype.startswith("I") and intent_job and any(
        kw in intent_job for kw in ("销售", "客服", "BD")
    ):
        alerts.append(
            f"你的 MBTI 偏内向（{mtype}），但意向岗位属于高强度人际类型，"
            "长期处在这样的环境会比较消耗能量，可以提前设计自己的'恢复机制'。"
        )

    if alerts:
        paragraphs.append(
            "**【潜在张力】**\n" + "\n".join(f"- {a}" for a in alerts[:2])
        )

    return "\n\n".join(paragraphs)


def _offline_job_summary(job: dict | None) -> str:
    if not job:
        return "尚未选择目标岗位，系统将使用推荐列表的第一位作为默认目标。"
    name = job.get("name") or "目标岗位"
    inds = "、".join((job.get("industries") or [])[:3]) or "多个行业"
    cities = "、".join((job.get("cities") or [])[:3]) or "主要一线城市"
    sr = job.get("salary_range") or [0, 0]
    salary = f"{sr[0]}-{sr[1]} 元/月" if sr[1] else "薪资待明确"
    skills = (job.get("skills") or [])[:6]
    skill_txt = "**" + "、".join(skills) + "**" if skills else "多项综合能力"
    edu = job.get("education") or "学历不限"
    raw_soft = job.get("soft_skills") or {}
    soft_skills = list(raw_soft.keys())[:4] if isinstance(raw_soft, dict) else list(raw_soft)[:4]

    lines: list[str] = []

    # 段1：市场概况
    lines.append(
        f"**【市场概况】** **{name}** 当前主要出现在 **{inds}** 领域，"
        f"招聘集中在 **{cities}**。"
        f"岗位的薪资区间约为 **{salary}**，学历要求一般为 **{edu}**。"
    )

    # 段2：核心能力要求
    lines.append(
        f"**【核心能力】** 该岗位最看重的硬技能是：{skill_txt}。"
    )
    if soft_skills:
        lines.append(
            f"同时，用人单位普遍要求具备 **{'、'.join(soft_skills)}** 等软素质，"
            "这些在面试中往往通过行为面试题（STAR 法则）来考察。"
        )

    # 段3：典型 JD 解读
    if job.get("description"):
        desc = str(job["description"]).strip().replace("\n", " ")
        lines.append(
            f"**【JD 解读】** 典型 JD 片段：{desc[:220]}…"
        )

    # 段4：工作节奏
    work_time = job.get("work_time") or ""
    benefits = job.get("benefits") or []
    if work_time or benefits:
        rhythm_parts = []
        if work_time:
            rhythm_parts.append(f"工作时间：{work_time}")
        if benefits:
            rhythm_parts.append(f"福利待遇：{'、'.join(benefits[:5])}")
        lines.append("**【工作节奏】** " + "。".join(rhythm_parts) + "。")

    # 段5：竞争力分析
    if sr[1]:
        if sr[1] >= 15000:
            lines.append(
                "**【竞争力分析】** 该岗位薪资处于中高水平，竞争相对激烈。"
                "建议你在简历中突出可量化的项目成果和差异化技能，"
                "同时准备 2-3 个深度案例以应对技术面或专业面。"
            )
        else:
            lines.append(
                "**【竞争力分析】** 该岗位薪资处于市场中位，入门门槛适中。"
                "重点在于展现你的学习能力和成长潜力，"
                "有相关实习或项目经历会显著提升竞争力。"
            )

    # 段6：总结建议
    lines.append(
        "**【对标建议】** 建议你结合自己的经历逐项对照这份岗位画像，"
        "找出已满足的条件和待补齐的差距，有针对性地在短期内提升匹配度。"
    )

    return "\n\n".join(lines)


def _offline_action_plan(prof: dict, target_job: dict | None) -> str:
    tname = (target_job or {}).get("name") or "目标岗位"
    my_skills = set(s.lower() for s in (prof.get("skills") or []))
    job_skills_all = (target_job or {}).get("skills") or []
    hit_skills = [s for s in job_skills_all if s.lower() in my_skills]
    miss_skills = [s for s in job_skills_all if s.lower() not in my_skills]

    intern = prof.get("internships") or []
    projects = prof.get("projects") or []
    total_exp = len(intern) + len(prof.get("work_experiences") or [])

    # 根据用户困惑模式注入针对性建议
    confusions = _detect_confusion(prof)
    wisdom_tips: list[str] = []
    for c in confusions[:2]:
        wisdom_tips.append(f"- {c['strategy']}")

    quotes = _pick_quotes(2)
    quote_line = " / ".join(f"「{q}」" for q in quotes) if quotes else ""

    parts: list[str] = [f"### 分阶段行动计划（面向 {tname}）\n"]

    # ── 张雪峰方法论注入 ──
    # 就业倒推法：从毕业后实际就业数据反推选择
    # 不可替代性检验：工资与不可替代性成正比
    # 中位数原则：关注中间 50% 的普遍去向
    zxf = _zxf_by_job(tname, 3)
    zxf_major = _zxf_by_major(
        (prof.get("education") or [{}])[0].get("major", "") if prof.get("education") else "",
        2,
    )
    # 去重
    seen_ids = set()
    zxf_all: list[dict] = []
    for q in zxf + zxf_major:
        if q["id"] not in seen_ids:
            seen_ids.add(q["id"])
            zxf_all.append(q)

    # ── 第 1 阶段：短期 ──
    parts.append("**第 1 阶段（1-4 周）：认知对齐**")
    parts.append(f"- 精读 3-5 份「{tname}」的真实 JD，标注高频出现的技能和经验要求")
    parts.append("- 用「就业倒推法」调研：这个岗位的中位数毕业生去了哪里、做了什么、拿多少薪资")
    parts.append("- 找 2-3 位从业者做 30 分钟职业访谈（学长/学姐/脉脉/LinkedIn）")
    if miss_skills:
        top_miss = miss_skills[:2]
        parts.append(f"- 启动第一项技能补强：**{top_miss[0]}**（选一门系统课程 + 动手项目）")
    parts.append("")

    # ── 第 2 阶段：中期 ──
    parts.append("**第 2 阶段（1-3 个月）：技能补齐 + 建立不可替代性**")
    if miss_skills:
        for i, skill in enumerate(miss_skills[:4], 1):
            parts.append(f"- 第 {i} 周重点：学习 **{skill}**，完成一个小型实践项目")
    else:
        parts.append("- 核心技能已较齐全，重点打磨深度：选 1-2 个技能做到「能讲 30 分钟」的水平")
    if hit_skills:
        parts.append(f"- 巩固已有优势（{'、'.join(hit_skills[:3])}），产出可展示的作品或案例")
    parts.append("- 思考你的「不可替代性」：哪些能力组合是别人难以复制的？把它作为简历的核心卖点")
    parts.append("- 在 GitHub / 掘金 / 个人博客沉淀可展示产出")
    parts.append("")

    # ── 第 3 阶段：实战 ──
    parts.append("**第 3 阶段（3-6 个月）：实战验证 + 500 强测试**")
    if total_exp == 0:
        parts.append(f"- 争取第一段与「{tname}」直接相关的实习（核心目标）")
    elif total_exp < 3:
        parts.append(f"- 再补充 1 段高质量实习，最好在目标行业内")
    else:
        parts.append("- 将最相关的一段经历提炼为 STAR 结构案例（情境→任务→行动→结果）")
    parts.append("- 完成 1 个中型项目并形成完整案例（问题→思路→方案→结果→反思）")
    parts.append("- 用「500 强测试」检验方向：看行业头部企业实际在招什么岗位、要求什么技能")
    parts.append("- 参加 1 次行业竞赛或开源协作，积累团队交付经验")
    parts.append("")

    # ── 里程碑检查点 ──
    parts.append("**里程碑检查点**")
    parts.append("- 第 1 个月末：能清晰描述「我为什么适合这个岗位」（1 分钟电梯演讲）")
    parts.append(f"- 第 3 个月末：{'、'.join(miss_skills[:3]) if miss_skills else '核心技能'}掌握到能做项目的水平")
    parts.append("- 第 6 个月末：简历上至少有 1 段强相关经历 + 2 个可展示项目")
    parts.append("- 用「10 年后压迫测试」自检：这个方向 10 年后还有竞争力吗？AI 会替代吗？")

    if wisdom_tips:
        parts.append("\n" + "\n".join(wisdom_tips))

    # 张雪峰语录
    if zxf_all:
        parts.append("\n**张雪峰说**")
        for q in zxf_all[:2]:
            parts.append(f"> 「{q['text']}」")
            if q.get("context"):
                parts.append(f"> *—— {q['context']}*")

    if quote_line:
        parts.append(f"\n{quote_line}")

    return "\n".join(parts)


def _collect_venn_data(
    prof: dict, assess: dict, target_job: dict
) -> dict:
    """把三类数据源按维度整理成可复用的结构。"""
    data: dict = {}

    # ===== 职业经验维度（来自 profile） =====
    skills = prof.get("skills") or []
    data["skills"] = skills
    data["top_skills"] = skills[:5]

    # 教育
    edu_items = [e for e in (prof.get("education") or []) if isinstance(e, dict)]
    best_edu = None
    if edu_items:
        best_edu = max(
            edu_items,
            key=lambda e: {"博士": 5, "硕士": 4, "本科": 3, "大专": 2, "高中": 1}.get(
                e.get("degree", ""), 0
            ),
        )
    data["school"] = (best_edu or {}).get("school", "")
    data["major"] = (best_edu or {}).get("major", "")
    data["degree"] = (best_edu or {}).get("degree", "")
    data["direction"] = (best_edu or {}).get("direction", "")

    # 实习 / 工作 / 项目
    internships = [e for e in (prof.get("internships") or []) if isinstance(e, dict)]
    works = [e for e in (prof.get("work_experiences") or []) if isinstance(e, dict)]
    projects = [e for e in (prof.get("projects") or []) if isinstance(e, dict)]
    data["internships"] = internships
    data["work_exps"] = works
    data["projects"] = projects

    def fmt_exp(e: dict) -> str:
        head = " · ".join(x for x in [e.get("company"), e.get("position")] if x)
        if e.get("desc"):
            head += "：" + e["desc"][:80]
        return head or ""

    data["exp_lines"] = [fmt_exp(e) for e in internships + works if fmt_exp(e)][:3]
    data["project_lines"] = [
        (p.get("name") or "") + (("：" + p["desc"][:60]) if p.get("desc") else "")
        for p in projects
    ][:3]

    # 软技能
    soft = prof.get("soft_skills") or {}
    if soft:
        top_soft = sorted(soft.items(), key=lambda x: -x[1])[:2]
        data["top_soft"] = [k for k, _ in top_soft]
    else:
        data["top_soft"] = []

    # 自我评价
    data["self_eval"] = (prof.get("self_eval") or "").strip()

    # ===== 个人兴趣维度（来自 intent + assessments） =====
    intent = prof.get("intent") or {}
    data["intent_industry"] = intent.get("industry", "")
    data["intent_job"] = intent.get("job", "")
    data["intent_city"] = intent.get("city", "")
    data["intent_salary"] = intent.get("salary_min") or 0
    data["intent_env"] = (intent.get("env") or "").strip()

    holland = assess.get("holland") or []
    if holland and len(holland) >= 6:
        hpairs = sorted(zip(HOLLAND_NAMES, holland), key=lambda x: -x[1])
        data["holland_code"] = _holland_code(holland)
        data["holland_top"] = hpairs[0][0]
        data["holland_top3"] = [n for n, _ in hpairs[:3]]
    else:
        data["holland_code"] = ""
        data["holland_top"] = ""
        data["holland_top3"] = []

    values = assess.get("values") or []
    if values and len(values) >= 6:
        vpairs = sorted(zip(VALUE_NAMES, values), key=lambda x: -x[1])
        data["value_top3"] = [n for n, _ in vpairs[:3]]
    else:
        data["value_top3"] = []

    mi = assess.get("multi_intel") or []
    if mi and len(mi) >= 6:
        mpairs = sorted(zip(MI_NAMES, mi), key=lambda x: -x[1])
        data["mi_top3"] = [n for n, _ in mpairs[:3]]
    else:
        data["mi_top3"] = []

    mbti = assess.get("mbti") or []
    data["mbti_type"] = _mbti_from_scores(mbti)
    data["mbti_label"] = MBTI_DESCRIPTIONS.get(data["mbti_type"], "")

    # ===== 岗位情况维度（来自 target_job） =====
    data["job_name"] = target_job.get("name", "")
    data["job_industries"] = (target_job.get("industries") or [])[:3]
    data["job_cities"] = (target_job.get("cities") or [])[:3]
    sr = target_job.get("salary_range") or [0, 0]
    data["job_salary"] = f"{sr[0]}-{sr[1]} 元/月" if sr[1] else ""
    data["job_edu_req"] = target_job.get("education", "")
    data["job_skills"] = (target_job.get("skills") or [])[:10]
    data["job_companies"] = (target_job.get("sample_companies") or [])[:3]

    # 技能命中与缺口
    student_set = {s.lower() for s in skills}
    hit = [s for s in (target_job.get("skills") or []) if s.lower() in student_set]
    miss = [s for s in (target_job.get("skills") or []) if s.lower() not in student_set]
    data["hit_skills"] = hit[:5]
    data["miss_skills"] = miss[:5]

    return data


def _offline_advise(
    zone: str,
    target_name: str,
    prof: dict,
    assess: dict | None = None,
    target_job: dict | None = None,
) -> str:
    """基于学生档案 + 测评 + 目标岗位，为韦恩四区生成具体建议。

    数据使用原则（严格分离）：
      - 基石区 = 经验 ∩ 兴趣   → 只引用经验数据 + 自我认识/意向数据
      - 稳定区 = 经验 ∩ 岗位   → 只引用经验数据 + 目标岗位数据
      - 前景区 = 兴趣 ∩ 岗位   → 只引用自我认识/意向数据 + 目标岗位数据
      - 发展区 = 三者交集      → 同时引用三类数据
    """
    d = _collect_venn_data(prof, assess or {}, target_job or {})

    # 智慧库注入
    confusions = _detect_confusion(prof, assess)
    def _wisdom_section() -> str:
        """根据用户困惑模式生成智慧提示段。"""
        parts: list[str] = []
        for c in confusions[:1]:
            parts.append(f"\n**💡 给你的提醒**：{c['strategy']}")
            if c.get("quotes"):
                parts.append(f"「{c['quotes'][0]}」")
        if not parts:
            qs = _pick_quotes(1)
            if qs:
                parts.append(f"\n**💡 送你一句话**：「{qs[0]}」")
        return "\n".join(parts)

    def _join_exp() -> str:
        """经验描述：学校+专业 + 第一段实习+项目 + 技能栈"""
        bits: list[str] = []
        if d["school"]:
            edu_bit = f"**{d['school']}**"
            if d["degree"]:
                edu_bit += f" {d['degree']}"
            if d["major"]:
                edu_bit += f"·{d['major']}"
            if d["direction"]:
                edu_bit += f"（{d['direction']}方向）"
            bits.append("就读于 " + edu_bit)
        if d["exp_lines"]:
            bits.append("实战经历包括 **" + "** / **".join(
                e.split("：")[0] for e in d["exp_lines"][:2]
            ) + "**")
        if d["project_lines"]:
            bits.append("项目产出涵盖 **" + "、".join(
                p.split("：")[0] for p in d["project_lines"][:2]
            ) + "**")
        if d["top_skills"]:
            bits.append("核心技能栈为 **" + "、".join(d["top_skills"]) + "**")
        if d["top_soft"]:
            bits.append("软技能突出项是 **" + "、".join(d["top_soft"]) + "**")
        return "，".join(bits) if bits else ""

    def _join_interest() -> str:
        """兴趣描述：Holland + MBTI + 价值观 + 优势智能 + 求职意向"""
        bits: list[str] = []
        if d["mbti_type"]:
            tag = d["mbti_label"].split("·")[0].strip() if d["mbti_label"] else ""
            bits.append(f"MBTI 为 **{d['mbti_type']}**" + (f"（{tag}）" if tag else ""))
        if d["holland_top"]:
            bits.append(f"主导霍兰德兴趣为 **{d['holland_top']}**（代码 **{d['holland_code']}**）")
        if d["value_top3"]:
            bits.append(f"核心价值观 **{' · '.join(d['value_top3'])}**")
        if d["mi_top3"]:
            bits.append(f"优势智能 **{' · '.join(d['mi_top3'])}**")
        if d["intent_job"] or d["intent_industry"]:
            aim_bits = []
            if d["intent_industry"]:
                aim_bits.append(d["intent_industry"])
            if d["intent_job"]:
                aim_bits.append(d["intent_job"])
            bits.append(f"明确的求职意向：**{' · '.join(aim_bits)}**")
        return "；".join(bits) if bits else ""

    def _join_job() -> str:
        """岗位描述：名称 + 行业 + 城市 + 薪资 + 核心技能 + 样本公司"""
        if not d["job_name"]:
            return ""
        bits: list[str] = [f"目标岗位 **{d['job_name']}**"]
        if d["job_industries"]:
            bits.append(f"主要出现在 **{' / '.join(d['job_industries'])}** 领域")
        if d["job_cities"]:
            bits.append(f"招聘集中在 **{' / '.join(d['job_cities'])}**")
        if d["job_salary"]:
            bits.append(f"薪资区间 **{d['job_salary']}**")
        if d["job_skills"]:
            bits.append(f"核心技能要求 **{'、'.join(d['job_skills'][:6])}**")
        if d["job_companies"]:
            bits.append(f"代表性公司有 **{'、'.join(d['job_companies'])}**")
        return "，".join(bits)

    # ========= 基石区：经验 ∩ 兴趣 =========
    if "基石" in zone:
        lines = [
            "**基石区 · 经验 ∩ 兴趣**\n",
            "这是你现在最容易获得成就感、最扎实的职业起点。"
            "它把你已经积累的能力，和你内心真正感兴趣的方向重合在一起。\n",
        ]
        exp_desc = _join_exp()
        int_desc = _join_interest()
        if exp_desc:
            lines.append(f"**【你的经验积累】** {exp_desc}。")
        else:
            lines.append("**【你的经验积累】** 你尚未在档案中填写完整的教育和实习信息，建议先去「我的档案」补齐。")
        if int_desc:
            lines.append(f"**【你的兴趣倾向】** {int_desc}。")
        else:
            lines.append("**【你的兴趣倾向】** 你尚未完成自我认识测评或求职意向，建议先去对应模块补齐。")

        lines.append("\n**建议行动：**")
        bullets: list[str] = []
        if d["top_skills"] and d["holland_top"]:
            bullets.append(
                f"围绕 **{d['top_skills'][0]}** 做一个能体现你 **{d['holland_top']}** 兴趣的小作品，"
                "这是最不费力又有成就感的起步方式。"
            )
        elif d["top_skills"]:
            bullets.append(
                f"围绕 **{d['top_skills'][0]}** 做出 1-2 个完整的小作品，累积可展示的产出。"
            )
        if d["exp_lines"]:
            first_exp_name = d["exp_lines"][0].split("：")[0]
            bullets.append(
                f"把 **{first_exp_name}** 的经历重新整理成故事：问题 → 方案 → 结果，作为简历的主线。"
            )
        if d["value_top3"]:
            bullets.append(
                f"你看重 **{d['value_top3'][0]}**，可以把它作为选择下一步机会的筛选器。"
            )
        if not bullets:
            bullets.append("先去档案和自我认识模块补齐数据，系统才能给出具体建议。")
        bullets.append(
            "**⚠ 注意**：基石区起步顺畅，但天花板可能较低，要把它当作跳板而非终点，"
            "同时为下一步的发展区积累能力。"
        )
        lines.append("\n".join(f"• {b}" for b in bullets))
        lines.append(_wisdom_section())
        return "\n".join(lines)

    # ========= 稳定区：经验 ∩ 岗位 =========
    if "稳定" in zone:
        lines = [
            f"**稳定区 · 经验 ∩ 岗位（{target_name}）**\n",
            "这是路径最清晰、风险最低的选择。"
            "它把你现有的能力，和目标岗位的硬性要求直接对齐。\n",
        ]
        exp_desc = _join_exp()
        job_desc = _join_job()
        if exp_desc:
            lines.append(f"**【你的经验积累】** {exp_desc}。")
        else:
            lines.append("**【你的经验积累】** 档案里还没有完整的经历数据，无法做精准对齐。")
        if job_desc:
            lines.append(f"**【目标岗位画像】** {job_desc}。")

        lines.append("\n**建议行动：**")
        bullets: list[str] = []
        if d["hit_skills"]:
            bullets.append(
                f"你的 **{'、'.join(d['hit_skills'])}** 正好命中该岗位的核心要求，"
                "在简历和面试中要重点突出这些关键词。"
            )
        if d["miss_skills"]:
            bullets.append(
                f"针对 {target_name} 需要补齐 **{'、'.join(d['miss_skills'])}**，"
                "每项安排 2-4 周的专项训练即可。"
            )
        if d["job_companies"]:
            bullets.append(
                f"优先投递 **{'、'.join(d['job_companies'][:3])}** 等代表性公司，"
                "同时关注它们的校招/社招通道。"
            )
        if d["exp_lines"]:
            first_exp_name = d["exp_lines"][0].split("：")[0]
            bullets.append(
                f"把 **{first_exp_name}** 的产出与 {target_name} 的 JD 要求对照，"
                "在简历里改写成岗位关心的动词（如'主导/优化/落地'）。"
            )
        if not bullets:
            bullets.append("先补齐档案里的实习、项目与技能，然后回来再看具体建议。")
        bullets.append(
            "**⚠ 注意**：稳定区安全，但缺少兴趣支撑，长期容易倦怠。"
            "建议每 6-12 个月做一次工作满意度自检，防止职业疲劳。"
        )
        lines.append("\n".join(f"• {b}" for b in bullets))
        lines.append(_wisdom_section())
        return "\n".join(lines)

    # ========= 前景区：兴趣 ∩ 岗位 =========
    if "前景" in zone:
        lines = [
            f"**前景区 · 兴趣 ∩ 岗位（{target_name}）**\n",
            "这是成长空间最大、最能激发长期热情的方向。"
            "它把你内心真正在乎的东西，和行业上升通道对齐。\n",
        ]
        int_desc = _join_interest()
        job_desc = _join_job()
        if int_desc:
            lines.append(f"**【你的兴趣倾向】** {int_desc}。")
        else:
            lines.append("**【你的兴趣倾向】** 请先在「自我认识」和「求职意向」模块补齐数据。")
        if job_desc:
            lines.append(f"**【目标岗位画像】** {job_desc}。")

        lines.append("\n**建议行动：**")
        bullets: list[str] = []

        # 兴趣 × 岗位的契合分析
        if d["holland_top"] and d["job_name"]:
            bullets.append(
                f"你的主导兴趣是 **{d['holland_top']}**，与 **{d['job_name']}** 的典型工作内容天然契合 —— "
                "相信这份直觉。"
            )
        if d["value_top3"]:
            bullets.append(
                f"你的核心价值观 **{' · '.join(d['value_top3'][:2])}**，"
                f"在 {target_name} 这类岗位上通常能被满足，属于"
                "'为自己工作'的状态。"
            )
        if d["miss_skills"]:
            bullets.append(
                f"**优先补齐门槛技能：{'、'.join(d['miss_skills'])}**。"
                "这些是进入该岗位的硬门槛，建议用 3-6 个月系统学习并做出 1-2 个作品。"
            )
        if d["job_industries"]:
            bullets.append(
                f"持续关注 **{' / '.join(d['job_industries'][:2])}** 领域的头部公司动态，"
                "订阅行业新闻、加入相关社群，积累行业视野。"
            )
        if not bullets:
            bullets.append("先去补齐测评和意向数据，系统才能给出精准的契合分析。")
        bullets.append(
            "**⚠ 注意**：前景区上限高但起步慢，要给自己 6-12 个月的爬坡期。"
            "不要在短期内和老手比薪资，要和过去的自己比成长。"
        )
        lines.append("\n".join(f"• {b}" for b in bullets))
        lines.append(_wisdom_section())
        return "\n".join(lines)

    # ========= 发展区：三者交集 =========
    if "发展" in zone:
        lines = [
            f"**发展区 · 三者交集（{target_name}）**\n",
            "这是你职业选择的最优落点。在这里，已有经验可直接转化为核心竞争力、"
            "个人兴趣能支撑长期深耕、行业前景又能持续提供上升通道 —— 三者相互强化、"
            "形成正向循环，是实现长期职业成长、走向稳定可持续发展的理想落点。\n",
        ]
        exp_desc = _join_exp()
        int_desc = _join_interest()
        job_desc = _join_job()
        if exp_desc:
            lines.append(f"**【经验】** {exp_desc}。")
        if int_desc:
            lines.append(f"**【兴趣】** {int_desc}。")
        if job_desc:
            lines.append(f"**【岗位】** {job_desc}。")

        lines.append("\n**建议行动：**")
        bullets: list[str] = []

        # 经验与岗位命中
        strength_parts: list[str] = []
        if d["hit_skills"]:
            strength_parts.append(f"**{'、'.join(d['hit_skills'])}**")
        if d["exp_lines"]:
            first = d["exp_lines"][0].split("：")[0]
            strength_parts.append(f"**{first}** 的经历")
        if strength_parts:
            bullets.append(
                f"你现在的核心竞争力是 {'，'.join(strength_parts)}，"
                f"这正是 {target_name} 最看重的东西，要放在简历最显眼的位置。"
            )

        # 兴趣与价值观驱动
        motivation_parts: list[str] = []
        if d["holland_top"]:
            motivation_parts.append(f"**{d['holland_top']}** 兴趣")
        if d["mbti_type"]:
            tag = d["mbti_label"].split("·")[0].strip() if d["mbti_label"] else ""
            motivation_parts.append(f"**{d['mbti_type']}" + (f"·{tag}" if tag else "") + "** 工作风格")
        if d["value_top3"]:
            motivation_parts.append(f"**{d['value_top3'][0]}** 的价值驱动")
        if motivation_parts:
            bullets.append(
                f"长期动力来源：{' + '.join(motivation_parts)}，"
                "它们会支撑你熬过爬坡期的低谷。"
            )

        # 智能优势对岗位的适配
        if d["mi_top3"] and d["job_name"]:
            bullets.append(
                f"你的优势智能 **{'、'.join(d['mi_top3'])}** 能让你在 {d['job_name']} 方向"
                "有天然的认知优势，应该刻意放大这些能力的使用场景。"
            )

        # 技能缺口补齐
        if d["miss_skills"]:
            bullets.append(
                f"下一步补齐：**{'、'.join(d['miss_skills'])}**。"
                "每项安排 4-6 周的专项练习，最后用一个真实作品或项目总结。"
            )
        else:
            bullets.append(
                "核心技能已基本齐全，下一步重点是把'能做到'升级为'做到了并能被看到'，"
                "主动输出案例、参与开源或社区分享。"
            )

        # 行业/城市锚点
        if d["job_cities"] and d["intent_city"]:
            if d["intent_city"] in " ".join(d["job_cities"]):
                bullets.append(
                    f"你的意向城市 **{d['intent_city']}** 正好是 {target_name} 的主要招聘城市之一，"
                    "机会密度高，可以放心聚焦。"
                )
            else:
                bullets.append(
                    f"你的意向城市 **{d['intent_city']}** 与 {target_name} 的主要招聘地 "
                    f"（{' / '.join(d['job_cities'][:2])}）不完全重合，"
                    "需要评估城市与岗位之间的取舍。"
                )

        lines.append("\n".join(f"• {b}" for b in bullets))

        # OKR 模板
        lines.append(f"\n**12 个月 OKR（面向 {target_name}）**\n")
        first_miss = d["miss_skills"][0] if d["miss_skills"] else "核心技能深度"
        lines.append(f"**Q1 目标**：补齐 **{first_miss}**，完成 1 个可展示的小项目")
        if len(d["miss_skills"]) > 1:
            lines.append(f"**Q2 目标**：掌握 **{d['miss_skills'][1]}**，获得 1 段相关实习")
        else:
            lines.append(f"**Q2 目标**：深化技能栈，获得 1 段相关实习经历")
        lines.append(f"**Q3 目标**：形成 2-3 个完整项目案例，开始投递 {target_name} 岗位")
        lines.append(f"**Q4 目标**：拿到满意的 Offer，或明确下一步调整方向")

        lines.append(_wisdom_section())
        return "\n".join(lines)

    return target_name


def _build_report(target_position_id: str | None, request: Request) -> dict:
    u = _current_user(request)
    s = store.get(u)
    prof = s.get("profile", {})
    assess = s.get("assessments") or {}
    recs = recommend(request, 5)

    # 1) target job: user-specified or Top 1 of recommendations
    target_job: dict | None = None
    if target_position_id:
        target_job = POS_BY_ID.get(target_position_id)
    if target_job is None and recs:
        target_job = POS_BY_ID.get(recs[0]["id"])

    target_name = target_job["name"] if target_job else "（未确定）"
    target_blob = _format_target_job(target_job) if target_job else "（未选择具体岗位）"
    experience_blob = _format_experiences(prof)
    intent_blob = _format_intent(prof)
    assess_blob = _format_assessments(assess)

    # 2-4) 三个摘要并行生成
    from concurrent.futures import ThreadPoolExecutor, as_completed

    def _gen_experience():
        return _ai_or(
            "你是资深职业规划顾问。请基于学生的在校经历、实习经历、工作经历和项目经历，"
            "用 150-250 字提炼他的职业经验摘要：长处在哪里？积累了哪些可迁移能力？"
            "存在哪些明显的经验缺口？语言要具体，不要泛泛而谈。",
            experience_blob,
            _offline_experience_summary(prof),
        )

    def _gen_interest():
        return _ai_or(
            "你是一位融合心理学背景的资深职业规划顾问。请针对学生的【求职意向】和【自我认识测评】"
            "做深度跨维度分析（300-450 字，Markdown 格式，**重点用加粗**），涵盖：\n"
            "1) **融合画像**：用一句话定义这个学生的职业档案类型（MBTI + 霍兰德 + 核心价值观 + 优势智能）\n"
            "2) **意向对齐**：分析学生的意向岗位/行业是否与测评结果一致。"
            "一致之处在哪？存在哪些潜在张力？是在用理性追兴趣，还是在用兴趣追理性？\n"
            "3) **认知优势 × 目标岗位**：优势智能组合能否支撑意向岗位？哪种智能是独特加分项？\n"
            "4) **价值观驱动力**：核心价值观 Top 3 如何影响他的择业底色？薪资期望与经济回报得分是否一致？\n"
            "5) **潜在冲突或盲区**：价值观是否存在对立高分、MBTI 与岗位类型是否错配、霍兰德与意向是否张力？\n"
            "6) **总结一句**：给出可操作的下一步关注点\n"
            "要求：语言温暖专业，有洞察而非罗列，避免术语堆砌。",
            f"【求职意向】\n{intent_blob}\n\n【自我认识测评】\n{assess_blob}",
            _offline_interest_summary(prof, assess),
        )

    def _gen_job():
        return _ai_or(
            "你是资深职业规划顾问。请基于以下岗位画像，为学生解读这个岗位：它做什么、"
            "用什么技能、典型职业路径、行业机会与挑战。150-250 字，不要堆砌术语，"
            "结尾一句提示学生此岗最看重什么素质。",
            target_blob,
            _offline_job_summary(target_job),
        )

    # 先提交 3 个摘要（后面 Venn 区域在 advise 函数定义后再提交）
    _pool = ThreadPoolExecutor(max_workers=8)
    f_exp = _pool.submit(_gen_experience)
    f_int = _pool.submit(_gen_interest)
    f_job = _pool.submit(_gen_job)

    # Four Venn zones — 基于韦恩图理论给出定制化建议
    VENN_THEORY = (
        "韦恩四区的理论含义：\n"
        "- 基石区 = 经验 ∩ 兴趣：依托已有经验、契合内在兴趣，是个体最易上手、"
        "最能获得成就感的领域，构成职业发展的扎实起点，但往往缺乏足够的长期上升空间。\n"
        "- 稳定区 = 经验 ∩ 岗位：贴合行业趋势、延续现有能力，路径清晰、风险较低，"
        "能提供稳定收入与职业安全感，但因缺少兴趣支撑，长期易陷入倦怠。\n"
        "- 前景区 = 兴趣 ∩ 岗位：兼具时代红利与内在热情，成长空间大、吸引力强，"
        "但因缺乏经验支撑，落地难度高、难以直接转化为现实竞争力。\n"
        "- 发展区 = 三者交集：经验可直接转化为核心竞争力，个人兴趣支撑长期深耕，"
        "行业前景提供持续上升通道，三者相互强化，是最理想的职业落点。"
    )

    # 每个区的数据使用范围 —— 严格划分
    ZONE_SCOPES = {
        "基石区": {
            "must_use": "【学生职业经验】和【学生自我认识+求职意向】",
            "must_not_use": "不要引用目标岗位的任何具体数据",
            "data": f"【学生职业经验】\n{experience_blob}\n\n"
                    f"【学生求职意向】\n{intent_blob}\n\n"
                    f"【学生自我认识测评】\n{assess_blob}",
        },
        "稳定区": {
            "must_use": "【学生职业经验】和【目标岗位画像】",
            "must_not_use": "不要涉及兴趣或测评结果",
            "data": f"【学生职业经验】\n{experience_blob}\n\n"
                    f"【目标岗位画像】\n{target_blob}",
        },
        "前景区": {
            "must_use": "【学生自我认识+求职意向】和【目标岗位画像】",
            "must_not_use": "不要涉及学生已有的经历或技能",
            "data": f"【学生求职意向】\n{intent_blob}\n\n"
                    f"【学生自我认识测评】\n{assess_blob}\n\n"
                    f"【目标岗位画像】\n{target_blob}",
        },
        "发展区": {
            "must_use": "【学生职业经验】、【学生自我认识+求职意向】和【目标岗位画像】",
            "must_not_use": "三类数据都要用到，寻找三者重合的落点",
            "data": f"【学生职业经验】\n{experience_blob}\n\n"
                    f"【学生求职意向】\n{intent_blob}\n\n"
                    f"【学生自我认识测评】\n{assess_blob}\n\n"
                    f"【目标岗位画像】\n{target_blob}",
        },
    }

    def advise(zone_key: str, zone_label: str) -> str:
        scope = ZONE_SCOPES[zone_key]
        return _ai_or(
            "你是一位资深职业规划顾问。请严格基于学生的真实档案数据，"
            f"为学生生成 '{zone_label}' 的专属建议（Markdown 格式，250-400 字）。\n\n"
            "**严格的数据使用规则**：\n"
            f"- 本区域**只能使用** {scope['must_use']} 这些数据源\n"
            f"- **禁止使用**：{scope['must_not_use']}\n\n"
            "**输出结构**：\n"
            "1) 开头一段：点明该区域对该学生的具体意义（≤80 字，不要照搬理论）\n"
            "2) 第二段：用 `**【...】**` 小标题明确标出你引用的数据源，"
            "并列出这些数据的关键事实（如学校/专业/技能/实习公司/Holland类型/价值观/目标岗位薪资等）\n"
            "3) 行动建议：3-4 条，每条必须至少引用 1 项上述事实，具体可执行，避免'你应该提升能力'这类废话\n"
            "4) 末尾一句：该区域的典型风险或注意事项\n\n"
            f"**理论参考**：\n{VENN_THEORY}\n\n"
            f"**目标岗位名称**：{target_name}",
            scope["data"],
            _offline_advise(zone_key, target_name, prof, assess, target_job),
        )

    # 4 个 Venn 区域 + Action plan 也提交到同一个线程池（8 个任务并行）
    f_foundation = _pool.submit(advise, "基石区", "基石区（经验 ∩ 兴趣）")
    f_stable = _pool.submit(advise, "稳定区", "稳定区（经验 ∩ 岗位）")
    f_prospect = _pool.submit(advise, "前景区", "前景区（兴趣 ∩ 岗位）")
    f_growth = _pool.submit(advise, "发展区", "发展区（三者交集，最优）")
    f_action = _pool.submit(
        _ai_or,
        "你是资深职业规划顾问。请针对该学生生成分阶段行动计划，Markdown 格式输出。"
        "包含：短期(1-3月)学习/实践、中期(3-12月)项目/实习、评估周期与指标。"
        "建议具体、可执行、不超过 600 字。",
        f"目标岗位：{target_name}\n"
        f"岗位画像：\n{target_blob}\n"
        f"学生经验：\n{experience_blob}\n"
        f"学生意向：\n{intent_blob}",
        _offline_action_plan(prof, target_job),
    )

    # 收集所有结果
    experience_sum = f_exp.result()
    interest_sum = f_int.result()
    job_sum = f_job.result()
    venn = {
        "foundation": f_foundation.result(),
        "stable": f_stable.result(),
        "prospect": f_prospect.result(),
        "growth": f_growth.result(),
    }
    action = f_action.result()
    _pool.shutdown(wait=False)

    report = {
        "summaries": {"experience": experience_sum, "interest": interest_sum, "job": job_sum},
        "venn": venn,
        "recommendations": recs,
        "action_plan": action,
        "target_position": (
            {"id": target_job["id"], "name": target_job["name"]} if target_job else None
        ),
    }
    store.patch("report", report, user=u)
    return report


@app.post("/api/report/generate")
def gen_report(request: Request, body: ReportIn | None = None):
    tid = body.target_position_id if body else None
    return _build_report(tid, request)


class PolishIn(BaseModel):
    text: str
    style: str = "专业"  # or 亲切


@app.get("/api/report/download")
def download_report(request: Request):
    """将已生成的报告导出为 Word 文档"""
    from docx import Document as DocxDocument
    from docx.shared import Pt, RGBColor
    from starlette.responses import Response

    u = _current_user(request)
    rep = store.get(u).get("report")
    if not rep:
        raise HTTPException(400, "请先生成报告")

    doc = DocxDocument()

    # 标题
    title = doc.add_heading("职业发展规划报告", level=0)
    target = rep.get("target_position") or {}
    if target.get("name"):
        doc.add_paragraph(f"目标岗位：{target['name']}")
    doc.add_paragraph(f"生成日期：{time.strftime('%Y-%m-%d')}")
    doc.add_paragraph()

    sums = rep.get("summaries") or {}

    # 一、职业经验总结
    doc.add_heading("一、职业经验总结", level=1)
    _docx_md(doc, sums.get("experience", ""))

    # 二、个人兴趣画像
    doc.add_heading("二、个人兴趣画像", level=1)
    _docx_md(doc, sums.get("interest", ""))

    # 三、目标岗位解读
    doc.add_heading("三、目标岗位解读", level=1)
    _docx_md(doc, sums.get("job", ""))

    # 四、韦恩四区分析
    venn = rep.get("venn") or {}
    zone_names = [
        ("foundation", "基石区（经验 ∩ 兴趣）"),
        ("stable", "稳定区（经验 ∩ 岗位）"),
        ("prospect", "前景区（兴趣 ∩ 岗位）"),
        ("growth", "发展区（三者交集）"),
    ]
    doc.add_heading("四、韦恩四区分析", level=1)
    for key, label in zone_names:
        doc.add_heading(label, level=2)
        _docx_md(doc, venn.get(key, ""))

    # 五、推荐岗位
    recs = rep.get("recommendations") or []
    if recs:
        doc.add_heading("五、推荐岗位", level=1)
        table = doc.add_table(rows=len(recs) + 1, cols=4)
        table.style = "Table Grid"
        for i, h in enumerate(["岗位", "行业", "城市", "匹配度"]):
            table.rows[0].cells[i].text = h
        for r, rec in enumerate(recs, 1):
            table.rows[r].cells[0].text = rec.get("name", "")
            table.rows[r].cells[1].text = "、".join(rec.get("industries", [])[:2])
            table.rows[r].cells[2].text = "、".join(rec.get("cities", [])[:2])
            table.rows[r].cells[3].text = f"{rec.get('match', 0)}%"

    # 六、行动计划
    doc.add_heading("六、行动计划", level=1)
    _docx_md(doc, rep.get("action_plan", ""))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=career_report.docx"},
    )


def _docx_md(doc, text: str):
    """简易 Markdown → Word 段落转换"""
    if not text:
        return
    from docx.shared import Pt, Cm
    num_counter = 0  # 手动计数有序列表
    for line in text.split("\n"):
        stripped = line.strip()
        if not stripped:
            num_counter = 0  # 空行重置编号
            continue
        # 标题
        if stripped.startswith("#### "):
            num_counter = 0
            doc.add_heading(stripped[5:].replace("**", ""), level=4)
        elif stripped.startswith("### "):
            num_counter = 0
            doc.add_heading(stripped[4:].replace("**", ""), level=3)
        elif stripped.startswith("## "):
            num_counter = 0
            doc.add_heading(stripped[3:].replace("**", ""), level=2)
        elif stripped.startswith("# "):
            num_counter = 0
            doc.add_heading(stripped[2:].replace("**", ""), level=1)
        elif stripped.startswith("- ") or stripped.startswith("* ") or stripped.startswith("• "):
            num_counter = 0
            p = doc.add_paragraph(style="List Bullet")
            _docx_bold_runs(p, stripped[2:])
        elif re.match(r"^\d+[\.\)、]\s*", stripped):
            num_counter += 1
            content = re.sub(r"^\d+[\.\)、]\s*", "", stripped)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            _docx_bold_runs(p, f"{num_counter}. {content}")
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(stripped[2:].replace("**", ""))
            run.italic = True
        elif stripped == "---":
            num_counter = 0
            continue
        else:
            num_counter = 0
            p = doc.add_paragraph()
            _docx_bold_runs(p, stripped)


def _docx_bold_runs(paragraph, text: str):
    """解析 **bold** 标记为 Word 加粗 run"""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


@app.post("/api/report/polish")
def polish(body: PolishIn):
    out = llm.chat(
        f"你是文字润色助手。请将以下内容按'{body.style}'风格润色，保持事实与结构。",
        body.text,
    )
    return {"polished": out}


class ReportPatchIn(BaseModel):
    path: str  # dot path within report, e.g. "venn.growth"
    value: Any


@app.post("/api/report/patch")
def patch_report(body: ReportPatchIn, request: Request):
    u = _current_user(request)
    rep = store.get(u).get("report") or {}
    parts = body.path.split(".")
    d = rep
    for p in parts[:-1]:
        d = d.setdefault(p, {})
    d[parts[-1]] = body.value
    store.patch("report", rep, user=u)
    return {"ok": True}


# ---------- production frontend ----------
# `npm run build` creates app/frontend/dist.  Serving it from FastAPI keeps the
# browser and API on one origin, which makes the project easy to deploy as one
# public web service and avoids CORS/proxy configuration in production.
FRONTEND_DIST = ROOT / "app" / "frontend" / "dist"
if FRONTEND_DIST.exists():
    from fastapi.staticfiles import StaticFiles
    from starlette.responses import FileResponse

    assets_dir = FRONTEND_DIST / "assets"
    if assets_dir.exists():
        app.mount("/assets", StaticFiles(directory=assets_dir), name="frontend-assets")

    @app.get("/{full_path:path}", include_in_schema=False)
    def frontend_spa(full_path: str):
        candidate = (FRONTEND_DIST / full_path).resolve()
        try:
            candidate.relative_to(FRONTEND_DIST.resolve())
        except ValueError:
            candidate = FRONTEND_DIST / "index.html"
        if candidate.is_file():
            return FileResponse(candidate)
        return FileResponse(FRONTEND_DIST / "index.html")
